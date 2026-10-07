"""Account-producing transfer; create-only, reviewed exact bytes, bounded work."""
import csv
import hashlib
import hmac
import io
import posixpath
import re
import zipfile
from decimal import Decimal, InvalidOperation
from defusedxml import ElementTree
from openpyxl.utils import get_column_letter
import json
import time
import uuid
from datetime import date

from openpyxl import Workbook
from pydantic import ValidationError
from sqlalchemy import func, or_, select
from app.core.config import get_settings
from app.db.models import User, MRProfile
from app.db.mr_models import MRDirectory
from app.db.zone_models import Zone
from app.db.headquarter_models import Headquarter
from app.schemas.mrs import MRFields
from app.services import mrs
from app.services.zone_transfer import workbook_rows, safe_text, UNSAFE
from app.services.zones import ZoneError

MAX_BYTES, MAX_ROWS, EXPORT_LIMIT = 2 * 1024 * 1024, 1000, 5000
COLUMNS = [
    ("employeeCode", "Employee Code"), ("name", "MR Name"), ("phone", "Phone No."),
    ("userId", "User ID"), ("email", "Email ID"), ("contactRequirement", "Contact Requirement"),
    ("hq", "HQ"), ("zoneId", "Assigned Zone"), ("dateOfJoining", "Date of Joining"),
    ("designation", "Designation"), ("reportingManagerId", "Reporting Manager"),
    ("paymentLimit", "Payment Limit"), ("doctorDaysLimit", "Doctor Days Limit"),
    ("status", "Status"), ("addressLine1", "Address Line 1"), ("addressLine2", "Address Line 2"),
    ("landmark", "Landmark"), ("pincode", "Pincode"), ("city", "City"), ("state", "State"),
    ("country", "Country"),
]
HEADERS = [title for _, title in COLUMNS]
AUDIT = ["Created By", "Created At", "Updated By", "Updated At"]
LEGACY = [title for title in HEADERS if title != "Contact Requirement"]


def invalid(message):
    raise mrs.MRError(message, 422, "mr_invalid_file")


def exact_workbook_rows(data):
    """Preserve monetary XML precision after the shared inert-workbook checks."""
    rows = workbook_rows(data, max_columns=25)
    if not rows or "Payment Limit" not in rows[0]:
        return rows
    column = rows[0].index("Payment Limit")
    letter = get_column_letter(column + 1)
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            book = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            sheet = next(node for node in book.iter() if node.tag.rsplit("}", 1)[-1] == "sheet")
            relation = next(value for key, value in sheet.attrib.items() if key.rsplit("}", 1)[-1] == "id")
            links = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            target = next(node.attrib["Target"] for node in links if node.attrib.get("Id") == relation)
            path = posixpath.normpath(target.lstrip("/") if target.startswith("/") else "xl/" + target)
            if not path.startswith("xl/worksheets/"):
                invalid("Unsupported worksheet location.")
            root = ElementTree.fromstring(archive.read(path))
            for cell in root.iter():
                if cell.tag.rsplit("}", 1)[-1] != "c" or cell.attrib.get("t", "n") != "n":
                    continue
                coordinate = re.fullmatch(r"([A-Z]{1,2})([1-9][0-9]{0,3})", cell.attrib.get("r", ""))
                if not coordinate or coordinate[1] != letter:
                    continue
                index = int(coordinate[2]) - 1
                value = next((child.text for child in cell if child.tag.rsplit("}", 1)[-1] == "v"), None)
                if not value or index < 1 or index >= len(rows) or column >= len(rows[index]):
                    continue
                amount = Decimal(value)
                rows[index][column] = ("invalid numeric payment" if not amount.is_finite() or amount < 0
                    or amount.adjusted() > 8 or amount.as_tuple().exponent < -2 else format(amount, "f"))
        return rows
    except (InvalidOperation, ValueError, KeyError, StopIteration, zipfile.BadZipFile):
        invalid("Malformed workbook payment or worksheet.")


def parse(data, filename):
    if not data or len(data) > MAX_BYTES:
        raise mrs.MRError("Use a non-empty CSV/XLSX no larger than 2 MiB.", 413, "mr_file_limit")
    if filename.lower().endswith(".csv"):
        try:
            rows = []
            for cells in csv.reader(io.StringIO(data.decode("utf-8-sig")), strict=True):
                if len(rows) >= MAX_ROWS + 1 or len(cells) > 25 or any(len(cell) > 10000 for cell in cells):
                    invalid("CSV exceeds 1,000 rows, 25 columns or 10,000 characters per cell.")
                rows.append(cells)
        except (UnicodeError, csv.Error):
            invalid("Use well-formed UTF-8 CSV.")
    elif filename.lower().endswith(".xlsx"):
        try:
            rows = exact_workbook_rows(data)
        except ZoneError as exc:
            invalid(exc.message)
    else:
        invalid("Only CSV and genuine XLSX are supported; no XLS/macros.")
    if not rows or rows[0] not in (HEADERS, LEGACY, HEADERS + AUDIT):
        invalid("Use the full 21-column MR schema, its 20-column legacy backup without Contact Requirement, or the current 25-column audit schema. Password, identity-link, version, deletion and unknown columns are forbidden. The reduced six-column mock template is not a backup.")
    header = rows[0]
    if len(rows) < 2:
        invalid("At least one MR is required.")
    parsed = []
    for number, cells in enumerate(rows[1:], 2):
        errors = [] if len(cells) == len(header) else ["Row width must match headers."]
        values = {}
        for key, title in COLUMNS:
            value = cells[header.index(title)].strip() if title in header and header.index(title) < len(cells) else ""
            if value.startswith("'") and (value[1:].startswith("'") or UNSAFE.match(value[1:])):
                value = value[1:]
            values[key] = value
        values["contactRequirement"] = (values["contactRequirement"] or "required").lower()
        values["status"] = values["status"].lower()
        if values["doctorDaysLimit"]:
            if values["doctorDaysLimit"].isdigit() and len(values["doctorDaysLimit"]) <= 4:
                values["doctorDaysLimit"] = int(values["doctorDaysLimit"])
            else:
                errors.append("Doctor Days Limit must be a whole number 0–3650.")
                values["doctorDaysLimit"] = 0
        parsed.append(dict(row=number, values=values, errors=errors))
    return parsed


def resolve_catalog(db, model, value, title):
    # Imported strings are labels/UUIDs belonging to this server catalogue only.
    candidates = []
    try:
        candidate = db.get(model, uuid.UUID(value))
        if candidate:
            candidates.append(candidate)
    except ValueError:
        pass
    candidates.extend(db.scalars(select(model).where(func.lower(model.name) == value.lower(), model.deleted_at.is_(None))))
    unique = {candidate.id: candidate for candidate in candidates}
    if len(unique) != 1:
        raise mrs.MRError(f"{title} '{title.lower()}' must resolve to one server record. Correct the label or supply its server UUID; local IDs cannot be imported.", 409, "mr_assignment")
    return next(iter(unique.values())).id


def resolve_manager(db, value, batch):
    if not value:
        return None
    value = value.lower()
    explicit = value.startswith(("user:", "employee:"))
    prefix, term = value.split(":", 1) if explicit else ("", value)
    choices = set()
    for record_id, fields in batch.items():
        if (prefix == "user" and fields["userId"].lower() == term or
            prefix == "employee" and fields["employeeCode"].lower() == term or
            not explicit and term in (str(record_id), fields["userId"].lower(), fields["employeeCode"].lower(), fields["name"].lower())):
            choices.add(record_id)
    query = select(MRDirectory.id).join(MRProfile, MRProfile.id == MRDirectory.id).join(User, User.id == MRProfile.user_id).where(MRDirectory.deleted_at.is_(None))
    if prefix == "user":
        query = query.where(func.lower(User.username) == term)
    elif prefix == "employee":
        query = query.where(func.lower(MRDirectory.employeeCode) == term)
    else:
        query = query.where(or_(func.lower(User.username) == term, func.lower(MRDirectory.employeeCode) == term,
                               func.lower(MRDirectory.name) == term))
        try:
            existing = db.get(MRDirectory, uuid.UUID(term))
            if existing and not existing.deleted_at:
                choices.add(existing.id)
        except ValueError:
            pass
    choices.update(db.scalars(query))
    if len(choices) != 1:
        raise mrs.MRError("Reporting Manager is missing or ambiguous. Use user:username or employee:code for one existing or same-batch MR.", 409, "mr_assignment")
    return next(iter(choices))


def review_state(db, parsed):
    # Stable synthetic planning IDs derive only from row positions + account names.
    batch = {uuid.uuid5(uuid.NAMESPACE_URL, "evexia-mr-import:" + str(row["row"]) + ":" + row["values"]["userId"].lower()): row["values"]
             for row in parsed}
    bodies, reviewed, identifiers, codes = {}, [], set(), set()
    for (record_id, source), row in zip(batch.items(), parsed):
        errors = list(row["errors"])
        values = dict(source)
        try:
            values["hq"] = resolve_catalog(db, Headquarter, values["hq"], "HQ")
            values["zoneId"] = resolve_catalog(db, Zone, values["zoneId"], "Assigned Zone")
            values["reportingManagerId"] = resolve_manager(db, values["reportingManagerId"], batch)
            body = MRFields(**values)
            mrs.identifiers_available(db, body)
            bodies[record_id] = body
            keys = {body.userId, body.email} - {""}
            if identifiers & keys:
                errors.append("Duplicate account identifier within this file.")
            if body.employeeCode.lower() in codes:
                errors.append("Duplicate employee code within this file.")
            identifiers.update(keys)
            codes.add(body.employeeCode.lower())
        except ValidationError as exc:
            names = sorted({str(e["loc"][0]) if e["loc"] else "contact rule" for e in exc.errors()})
            errors.append("Invalid business fields: " + ", ".join(names) + ". Use the documented bounds.")
        except mrs.MRError as exc:
            errors.append(exc.message)
        reviewed.append(dict(row=row["row"], name=source["name"], userId=source["userId"], errors=errors))
    for record_id, body in bodies.items():
        position = list(batch).index(record_id)
        try:
            standin = type("Existing", (), {"id": record_id, "hq": None, "zoneId": None, "reportingManagerId": None})()
            mrs.validate_assignments(db, body, standin, bodies)
        except mrs.MRError as exc:
            reviewed[position]["errors"].append(exc.message)
    snapshot = []
    for record_id, body in bodies.items():
        # Assignment versions bind review to the actual resolved server catalogue.
        refs = []
        for model, target in ((Headquarter, body.hq), (Zone, body.zoneId), (MRDirectory, body.reportingManagerId)):
            if target and target not in bodies:
                entry = db.get(model, target)
                refs.append((str(target), entry.version if entry else None))
        snapshot.append((str(record_id), body.model_dump(mode="json"), refs))
    return reviewed, bodies, snapshot


def digest_for(actor, data, filename, snapshot):
    message = json.dumps(["mr", str(actor.user.id), actor.session_id, actor.user.identity_version,
                          filename, hashlib.sha256(data).hexdigest(), snapshot], sort_keys=True).encode()
    return hmac.new(get_settings().signing_key.encode(), message, hashlib.sha256).hexdigest()


def transfer(db, actor, data, filename, confirm=False, digest=None):
    parsed = parse(data, filename)
    def review():
        current = mrs.authorize(db, actor, provision=True)
        mrs.graph_lock(db)
        rows, bodies, snapshot = review_state(db, parsed)
        fingerprint = digest_for(current, data, filename, snapshot)
        if confirm and not hmac.compare_digest(digest or "", fingerprint):
            raise mrs.MRError("File, identity, session or resolved assignments changed. Review again.", 409, "mr_review_changed")
        valid = not any(row["errors"] for row in rows)
        if confirm and not valid:
            raise mrs.MRError("Invalid rows or new conflicts. Nothing imported. Review again.", 409, "mr_import_conflict")
        db.commit()
        return dict(rows=rows, valid=valid, digest=fingerprint), bodies
    response, planned = mrs.transaction(db, review)
    if not confirm:
        return response
    # Password work is outside row transactions; failures publish no accounts.
    prepared = mrs.prepare_passwords(db, actor, [None] * len(planned))
    def commit():
        started = time.monotonic()
        current = mrs.authorize(db, actor, provision=True)
        mrs.graph_lock(db)
        rows, bodies, snapshot = review_state(db, parsed)
        if any(row["errors"] for row in rows) or not hmac.compare_digest(digest or "", digest_for(current, data, filename, snapshot)):
            raise mrs.MRError("Conflicts changed during credential preparation. Nothing imported. Review again.", 409, "mr_import_conflict")
        # Insert users/profiles/rows with manager links temporarily empty, then
        # wire the reviewed graph atomically (same-batch parents may come later).
        links, credentials = [], []
        for (record_id, body), password in zip(bodies.items(), prepared):
            if time.monotonic() - started > 20:
                raise mrs.MRError("Import transaction exceeded 20 seconds. Nothing was imported. Split the file and review again.",
                                  408, "mr_transaction_timeout")
            parent = body.reportingManagerId
            row = mrs.insert(db, current, body.model_copy(update={"reportingManagerId": None}), password, record_id)
            links.append((row, parent))
            credentials.append(dict(userId=body.userId, password=password[0]))
        for row, parent in links:
            row.reportingManagerId = parent
        db.flush()
        db.commit()
        return dict(imported=len(credentials), credentials=credentials)
    return mrs.transaction(db, commit)


def encode(rows, format, audit=False):
    header = HEADERS + (AUDIT if audit else [])
    if format == "csv":
        stream = io.StringIO()
        writer = csv.writer(stream, lineterminator="\r\n")
        writer.writerow(header)
        writer.writerows([[safe_text(value) for value in row] for row in rows])
        return ("\ufeff" + stream.getvalue()).encode("utf-8")
    book = Workbook()
    sheet = book.active
    sheet.title = "MR Master"
    sheet.append(header)
    for values in rows:
        sheet.append([safe_text(value) for value in values])
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def sample(format):
    # Explicit fictional labels are not silently created: operators replace them.
    values = dict(employeeCode="MR-001", name="Example MR", phone="", userId="example.mr",
                  email="", contactRequirement="optional", hq="Replace with active HQ",
                  zoneId="Replace with active Zone", dateOfJoining=date.today().isoformat(),
                  designation="Medical Representative", reportingManagerId="", paymentLimit="0.00",
                  doctorDaysLimit="0", status="inactive", addressLine1="Example address", addressLine2="",
                  landmark="Example landmark", pincode="110001", city="Delhi", state="Delhi", country="India")
    return encode([[values[key] for key, _ in COLUMNS]], format)


def export(db, actor, query, status, zone_id, hq_id, format):
    def work():
        mrs.authorize(db, actor, lock=False)
        clauses = mrs.predicates(query, status, zone_id, hq_id)
        count = db.scalar(select(func.count()).select_from(MRDirectory).where(*clauses))
        if count > EXPORT_LIMIT:
            raise mrs.MRError("More than 5,000 MRs match. Narrow the filters; nothing downloaded.", 409, "mr_export_limit")
        records = list(db.scalars(select(MRDirectory).where(*clauses).order_by(MRDirectory.created_at.desc(), MRDirectory.id.desc()).limit(EXPORT_LIMIT + 1)))
        if len(records) > EXPORT_LIMIT:
            raise mrs.MRError("More than 5,000 MRs match. Narrow the filters.", 409, "mr_export_limit")
        rows = []
        for record in records:
            values = mrs.projection(db, record)
            values["hq"], values["zoneId"] = values["hqName"], values["zoneName"]
            if record.reportingManagerId:
                manager_profile = db.get(MRProfile, record.reportingManagerId)
                values["reportingManagerId"] = "user:" + db.get(User, manager_profile.user_id).username
            else:
                values["reportingManagerId"] = ""
            rows.append([str(values[key]) for key, _ in COLUMNS] +
                        ([str(values[key]) for key in ("createdBy", "createdAt", "updatedBy", "updatedAt")] if format == "xlsx" else []))
        db.commit()
        return encode(rows, format, audit=format == "xlsx")
    return mrs.transaction(db, work)
