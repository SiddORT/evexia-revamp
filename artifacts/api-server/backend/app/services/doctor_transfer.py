"""Bounded create-only Doctor transfers, identity/file/reference bound review."""
import csv
import hashlib
import hmac
import io
import json
import posixpath
import re
import zipfile
from decimal import Decimal, InvalidOperation
from defusedxml import ElementTree
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from pydantic import ValidationError
from sqlalchemy import func, select
from app.core.config import get_settings
from app.db.doctor_models import DoctorDirectory as Doctor
from app.db.mr_models import MRDirectory
from app.db.models import MRProfile, User
from app.db.zone_models import Zone
from app.schemas.doctors import DoctorFields
from app.services import doctors, mrs
from app.services.mr_transfer import resolve_manager
from app.services.zone_transfer import workbook_rows, safe_text, UNSAFE
from app.services.zones import ZoneError

MAX_BYTES, MAX_ROWS, EXPORT_LIMIT = 2 * 1024 * 1024, 1000, 5000
COLUMNS = [
    ("name", "Doctor Name"), ("phone", "Phone"), ("dialCode", "Dial Code"), ("dialCountry", "Dial Country"),
    ("alternatePhone", "Alternate Phone"), ("email", "Email"), ("contactRequirement", "Contact Requirement"),
    ("dateOfJoining", "Date of Joining"), ("registrationNumber", "Registration Number"),
    ("qualification", "Qualification"), ("clinicName", "Clinic Name"), ("mrId", "Assigned MR"), ("zoneName", "Zone"),
    ("invoiceType", "Invoice Type"), ("gstNumber", "GST Number"), ("drugLicenceNumber", "Drug Licence Number"),
    ("orderDiscount", "Order Discount"), ("daysLimit", "Days Limit"), ("paymentLimit", "Payment Limit"),
    ("status", "Status"), ("verification", "Verification"), ("addressLine1", "Address Line 1"),
    ("addressLine2", "Address Line 2"), ("landmark", "Landmark"), ("pincode", "Pincode"),
    ("country", "Country"), ("state", "State"), ("city", "City"),
]
HEADERS = [title for _, title in COLUMNS]
CURRENT_COLUMNS = COLUMNS + [("mrName", "MR Name")]
CURRENT_HEADERS = HEADERS + ["MR Name"]
AUDIT = ["Created By", "Created At", "Updated By", "Updated At"]
LEGACY = [title for title in HEADERS if title != "Contact Requirement"]
from app.schemas.phone import DIAL


def invalid(message):
    raise doctors.DoctorError(message, 422, "doctor_invalid_file")


def exact_workbook_rows(data):
    rows = workbook_rows(data, max_columns=33)
    if not rows:
        return rows
    monetary = {get_column_letter(i + 1): key for i, key in enumerate(rows[0]) if key in ("Payment Limit", "Order Discount")}
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            book = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            sheet = next(n for n in book.iter() if n.tag.rsplit("}", 1)[-1] == "sheet")
            rel = next(v for k, v in sheet.attrib.items() if k.rsplit("}", 1)[-1] == "id")
            links = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            target = next(n.attrib["Target"] for n in links if n.attrib.get("Id") == rel)
            path = posixpath.normpath(target.lstrip("/") if target.startswith("/") else "xl/" + target)
            if not path.startswith("xl/worksheets/"):
                invalid("Unsupported worksheet location.")
            root = ElementTree.fromstring(archive.read(path))
            for cell in root.iter():
                coordinate = re.fullmatch(r"([A-Z]{1,2})([1-9][0-9]{0,3})", cell.attrib.get("r", ""))
                if cell.tag.rsplit("}", 1)[-1] != "c" or cell.attrib.get("t", "n") != "n" or not coordinate or coordinate[1] not in monetary:
                    continue
                index = int(coordinate[2]) - 1
                column = rows[0].index(monetary[coordinate[1]])
                value = next((n.text for n in cell if n.tag.rsplit("}", 1)[-1] == "v"), None)
                if not value or index < 1 or index >= len(rows) or column >= len(rows[index]):
                    continue
                amount = Decimal(value)
                rows[index][column] = ("invalid decimal" if not amount.is_finite() or amount < 0 or amount.adjusted() > 12
                                       or amount.as_tuple().exponent < -2 else format(amount, "f"))
        return rows
    except (ValueError, InvalidOperation, KeyError, StopIteration, zipfile.BadZipFile):
        invalid("Malformed workbook decimal or worksheet.")


def parse(data, filename):
    if not data or len(data) > MAX_BYTES:
        raise doctors.DoctorError("Use a non-empty CSV/XLSX no larger than 2 MiB.", 413, "doctor_file_limit")
    if filename.lower().endswith(".csv"):
        try:
            rows = []
            for cells in csv.reader(io.StringIO(data.decode("utf-8-sig")), strict=True):
                if len(rows) >= MAX_ROWS + 1 or len(cells) > 33 or any(len(v) > 10000 for v in cells):
                    invalid("File exceeds 1,000 records, 33 columns or 10,000 characters per cell.")
                rows.append(cells)
        except (UnicodeError, csv.Error):
            invalid("Use well-formed UTF-8 CSV.")
    elif filename.lower().endswith(".xlsx"):
        try:
            rows = exact_workbook_rows(data)
        except ZoneError as exc:
            invalid(exc.message)
    else:
        invalid("Only UTF-8 CSV and genuine XLSX are supported.")
    if not rows or rows[0] not in (HEADERS, HEADERS + AUDIT, LEGACY, LEGACY + AUDIT, CURRENT_HEADERS, CURRENT_HEADERS + AUDIT):
        invalid("Use the full Doctor sample/export columns in their original order. Older full CSVs without Contact Requirement default to optional. Reduced mock templates, credentials, local IDs and unknown columns are not supported.")
    if len(rows) < 2:
        invalid("At least one doctor is required.")
    parsed = []
    for number, cells in enumerate(rows[1:], 2):
        errors = [] if len(cells) == len(rows[0]) else ["Row width must match headers."]
        values = {}
        for key, title in COLUMNS:
            index = rows[0].index(title) if title in rows[0] else -1
            value = cells[index].strip() if 0 <= index < len(cells) else ""
            if value.startswith("'") and (value[1:].startswith("'") or UNSAFE.match(value[1:])):
                value = value[1:]
            elif key in ("phone", "alternatePhone", "registrationNumber", "pincode", "gstNumber", "drugLicenceNumber") and re.fullmatch(r"'\d+", value):
                value = value[1:]
            values[key] = value
        values["contactRequirement"] = (values["contactRequirement"] or "optional").lower()
        values["verification"] = (values["verification"] or "unverified").lower()
        for key in ("status", "invoiceType"):
            values[key] = values[key].lower()
        values["dialCountry"] = values["dialCountry"].upper() or "IN"
        for key in ("orderDiscount", "paymentLimit"):
            values[key] = values[key] or "0"
        days = values["daysLimit"] or "0"
        if not days.isdigit() or len(days) > 10:
            errors.append("Days Limit must be a bounded nonnegative whole number.")
            days = "0"
        values["daysLimit"] = int(days)
        parsed.append(dict(row=number, values=values, errors=errors))
    return parsed


def review_state(db, parsed):
    reviewed, bodies, snapshot, seen = [], [], [], set()
    for row in parsed:
        values = dict(row["values"])
        errors = list(row["errors"])
        try:
            # Exact UUID, user:username, employee:code or unique server name.
            # No browser-local ID conversion or fuzzy name mapping.
            mr_id = resolve_manager(db, values["mrId"], {})
            if not mr_id:
                raise doctors.DoctorError("Assigned MR is required.", 409, "doctor_assignment")
            values["mrId"] = mr_id
            supplied_zone = values.pop("zoneName")
            supplied_dial = values.pop("dialCode")
            verification = values.pop("verification")
            if verification not in ("verified", "unverified"):
                errors.append("Verification must be verified or unverified.")
            if supplied_dial and supplied_dial != DIAL.get(values["dialCountry"]):
                errors.append("Dial Code disagrees with Dial Country.")
            body = DoctorFields(**values)
            mr, zone = doctors.assignment(db, body.mrId)
            if supplied_zone and supplied_zone.lower() != zone.name.lower():
                errors.append("Zone does not match the assigned server MR. Zone is derived, not writable.")
            doctors.unique(db, body.registrationNumber)
            key = body.registrationNumber.lower()
            if key in seen:
                errors.append("Duplicate registration number within this file.")
            seen.add(key)
            bodies.append((body, verification))
            snapshot.append((body.model_dump(mode="json"), verification, str(mr.id), mr.version, str(zone.id), zone.version))
        except ValidationError as exc:
            errors.extend(["Invalid " + ".".join(map(str, e["loc"])) + ": " + e["msg"] for e in exc.errors()[:10]])
        except mrs.MRError as exc:
            errors.append(exc.message)
        reviewed.append(dict(row=row["row"], name=row["values"]["name"],
                             registrationNumber=row["values"]["registrationNumber"], errors=errors))
    return reviewed, bodies, snapshot


def digest_for(actor, data, filename, snapshot):
    message = json.dumps(["doctor", str(actor.user.id), actor.session_id, actor.user.identity_version,
                          filename, hashlib.sha256(data).hexdigest(), snapshot], sort_keys=True).encode()
    return hmac.new(get_settings().signing_key.encode(), message, hashlib.sha256).hexdigest()


def transfer(db, actor, data, filename, confirm=False, digest=None):
    parsed = parse(data, filename)
    def work():
        current = doctors.authorize(db, actor, "import")
        mrs.graph_lock(db)
        rows, bodies, snapshot = review_state(db, parsed)
        fingerprint = digest_for(current, data, filename, snapshot)
        valid = not any(row["errors"] for row in rows)
        if confirm and (not valid or not hmac.compare_digest(digest or "", fingerprint)):
            raise doctors.DoctorError("File, session, assignments or conflicts changed. Nothing imported. Review again.", 409, "doctor_review_changed")
        if confirm:
            for body, verification in bodies:
                doctors.insert(db, current, body, verification)
            result = dict(imported=len(bodies))
        else:
            result = dict(rows=rows, valid=valid, digest=fingerprint)
        db.commit()
        return result
    return doctors.transaction(db, work)


def encode(rows, format, audit=False):
    headers = (CURRENT_HEADERS if audit else HEADERS) + (AUDIT if audit else [])
    if format == "csv":
        stream = io.StringIO()
        writer = csv.writer(stream, lineterminator="\r\n")
        writer.writerow(headers)
        text_keys = {"phone", "alternatePhone", "registrationNumber", "pincode", "gstNumber", "drugLicenceNumber"}
        writer.writerows([[
            "'" + str(value) if index < len(COLUMNS) and COLUMNS[index][0] in text_keys and re.fullmatch(r"\d+", str(value))
            else safe_text(value) for index, value in enumerate(row)
        ] for row in rows])
        return ("\ufeff" + stream.getvalue()).encode()
    book = Workbook()
    sheet = book.active
    sheet.title = "Doctor Master"
    sheet.append(headers)
    for values in rows:
        sheet.append([safe_text(value) for value in values])
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def sample(format):
    values = {key: "" for key, _ in COLUMNS}
    values.update(name="Example Doctor", dialCountry="IN", dialCode="+91", contactRequirement="optional",
                  registrationNumber="EXAMPLE-001", qualification="MBBS", mrId="user:replace.with.server.mr",
                  invoiceType="normal", orderDiscount="0.00", daysLimit="0", paymentLimit="0.00",
                  status="active", verification="unverified", addressLine1="Example Clinic Road",
                  landmark="Example landmark", pincode="110001", country="India", state="Delhi", city="Delhi")
    return encode([[values[key] for key, _ in COLUMNS]], format)


def export(db, actor, query="", status="all", zone_id="", mr_id="", state="", format="csv"):
    def work():
        doctors.authorize(db, actor, "export", lock=False)
        clauses = doctors.predicates("", status, zone_id, mr_id, state)
        count = db.scalar(select(func.count()).select_from(Doctor).where(*clauses))
        if not query.strip() and count > EXPORT_LIMIT:
            raise doctors.DoctorError("More than 5,000 doctors match. Narrow the filters.", 409, "doctor_export_limit")
        rows = doctors.encrypted.complete(db, Doctor, clauses, None,
                                         prepare=lambda rows: doctors.search_matcher(db, query, rows))
        if len(rows) > EXPORT_LIMIT:
            raise doctors.DoctorError("More than 5,000 doctors match. Narrow the filters.", 409, "doctor_export_limit")
        result = []
        from app.services.patients import doctor_context
        context = doctor_context(db, [row.id for row in rows])
        for row in rows:
            values = doctors.projection(db, row, context)
            mr = db.get(MRDirectory, row.mrId)
            profile = db.get(MRProfile, row.mrId)
            # Explicit durable server reference, not a potentially ambiguous display label.
            values["mrId"] = "user:" + db.get(User, profile.user_id).username if mr and not mr.deleted_at and profile else str(row.mrId)
            values["dialCode"] = DIAL[row.dialCountry]
            result.append([str(values[key]) if values[key] is not None else "" for key, _ in CURRENT_COLUMNS] +
                          [str(values[key]) for key in ("createdBy", "createdAt", "updatedBy", "updatedAt")])
        db.commit()
        return encode(result, format, True)
    return doctors.transaction(db, work)
