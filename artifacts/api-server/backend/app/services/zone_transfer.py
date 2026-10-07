"""Bounded, inert spreadsheet transfer. File bytes are never persisted."""
import csv
import hashlib
import io
import re
import zipfile

from defusedxml import ElementTree
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy import Text, cast, func, select
from sqlalchemy.dialects.postgresql import ARRAY

from app.db.zone_models import Zone
from app.schemas.zones import ZoneFields
from app.services import zones

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1000
EXPORT_LIMIT = 1000
HEADERS = ["Zone Name", "Status", "Created By", "Created At", "Updated By", "Updated At"]
UNSAFE = re.compile(r"^[\s\x00-\x1f]*[=+\-@]")


def safe_text(value):
    text = str(value)
    return "'" + text if UNSAFE.match(text) or text.startswith("'") else text


def invalid(message="Invalid file. Use UTF-8 CSV or a genuine, single-sheet .xlsx template."):
    raise zones.ZoneError(message, 422, "zone_invalid_file")


def workbook_rows(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 100 or sum(item.file_size for item in members) > 8 * 1024 * 1024:
                invalid("Workbook exceeds decompressed size or archive-entry limits.")
            names = [item.filename for item in members]
            if len(set(names)) != len(names):
                invalid()
            cells = 0
            for item in members:
                if item.flag_bits & 1 or item.file_size > 4 * 1024 * 1024:
                    invalid()
                name = item.filename.lower()
                if any(part in name for part in ("vbaproject", "externallink", "macros", "embedding")):
                    invalid("Macros, embedded objects and external spreadsheet links are not supported.")
                if name.endswith((".xml", ".rels")):
                    root = ElementTree.fromstring(archive.read(item))
                    for node in root.iter():
                        tag = node.tag.rsplit("}", 1)[-1]
                        if tag == "f" or node.attrib.get("TargetMode") == "External":
                            invalid("Formulas and external links are not supported.")
                        if tag == "Override" and "macroEnabled" in node.attrib.get("ContentType", ""):
                            invalid("Macro-enabled workbooks are not supported.")
                        if tag == "c":
                            cells += 1
                            if cells > 6006:
                                invalid("Workbook exceeds the 6,006-cell limit.")
                            if name.startswith("xl/worksheets/"):
                                coordinate = re.fullmatch(r"([A-F])([1-9][0-9]{0,3})", node.attrib.get("r", ""))
                                if not coordinate or int(coordinate[2]) > MAX_ROWS + 1:
                                    invalid("Workbook cell coordinates exceed the row/column bounds.")
                        if tag == "row" and name.startswith("xl/worksheets/"):
                            number = node.attrib.get("r", "")
                            if not number.isdigit() or not 1 <= int(number) <= MAX_ROWS + 1:
                                invalid("Workbook row coordinates exceed the limit.")
                        if node.text and len(node.text) > 10000:
                            invalid("Workbook cell or XML text exceeds the limit.")
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.worksheets) != 1:
                invalid("Use exactly one worksheet.")
            sheet = book.worksheets[0]
            if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > 6:
                invalid("Workbook exceeds 1,000 records or six columns.")
            # Do not trust an understated producer-controlled dimension to hide
            # rows/cells. Coordinates were bounded independently above.
            sheet.reset_dimensions()
            rows = []
            for row in sheet.iter_rows():
                if len(rows) >= MAX_ROWS + 1:
                    invalid("Import is limited to 1,000 records.")
                values = []
                for cell in row:
                    if cell.data_type == "f":
                        invalid("Formulas are not supported.")
                    value = "" if cell.value is None else str(cell.value)
                    if len(value) > 10000:
                        invalid("A cell exceeds the limit.")
                    values.append(value)
                rows.append(values)
            if rows:
                width = len(rows[0])
                rows = [row + [""] * max(0, width - len(row)) for row in rows]
            return rows
        finally:
            book.close()
    except zones.ZoneError:
        raise
    except Exception:
        invalid()


def parse(data, filename):
    if not data or len(data) > MAX_BYTES:
        raise zones.ZoneError("Use a non-empty file no larger than 2 MiB.", 413, "zone_file_limit")
    if filename.lower().endswith(".csv"):
        try:
            reader = csv.reader(io.StringIO(data.decode("utf-8-sig")), strict=True)
            rows = []
            for row in reader:
                if len(rows) >= MAX_ROWS + 1 or len(row) > 6 or any(len(cell) > 10000 for cell in row):
                    invalid("CSV exceeds 1,000 records, six columns or cell limits.")
                rows.append(row)
        except (UnicodeError, csv.Error):
            invalid("CSV must be valid UTF-8 with well-formed quoted cells.")
    elif filename.lower().endswith(".xlsx"):
        rows = workbook_rows(data)
    else:
        invalid("Only .csv and .xlsx are supported; .xls and macro-enabled files are not accepted.")
    if not rows:
        invalid("A header and at least one record are required.")
    headers = [cell.strip() for cell in rows[0]]
    if headers not in (HEADERS[:2], HEADERS):
        invalid("Use Zone Name/Status, or the six-column Zone backup schema.")
    reviewed, seen = [], set()
    for number, cells in enumerate(rows[1:], 2):
        errors = []
        if len(cells) != len(headers):
            errors.append("Row must match the header column count.")
        name = cells[0].strip() if cells else ""
        # Our six-column portable exports escape leading apostrophes as well as
        # formula-like text. Decode only this documented export convention.
        if len(headers) == 6 and name.startswith("'") and (name[1:].startswith("'") or UNSAFE.match(name[1:])):
            name = name[1:]
        status = cells[1].strip().lower() if len(cells) > 1 else ""
        try:
            ZoneFields(name=name, status=status)
        except ValidationError:
            errors.append("Name must be 1–200 characters and status Active or Inactive.")
        key = name.lower()
        if key in seen:
            errors.append("Duplicate name within this file.")
        seen.add(key)
        reviewed.append(dict(row=number, name=name, status=status, errors=errors))
    if not reviewed:
        invalid("At least one record is required.")
    return reviewed


def transfer(db, actor, data, filename, confirm=False, digest=None):
    rows = parse(data, filename)
    fingerprint = hashlib.sha256(data).hexdigest()
    if confirm and digest != fingerprint:
        raise zones.ZoneError("The file changed. Review it again before confirming.", 409, "zone_review_changed")
    def work():
        current = zones.authorize(db, actor)
        # Match exactly the database lower() uniqueness semantics.
        names = [row["name"] for row in rows]
        normalized = list(db.scalars(select(func.lower(func.unnest(cast(names, ARRAY(Text)))))))
        seen = set()
        existing = set(db.scalars(select(func.lower(Zone.name)).where(
            Zone.deleted_at.is_(None), func.lower(Zone.name).in_(normalized))))
        for row, key in zip(rows, normalized):
            if key in seen and "Duplicate name within this file." not in row["errors"]:
                row["errors"].append("Duplicate name within this file.")
            if key in existing:
                row["errors"].append("A non-deleted zone already uses this name.")
            seen.add(key)
        valid = not any(row["errors"] for row in rows)
        if confirm:
            if not valid:
                raise zones.ZoneError("Import conflicts or invalid rows. Nothing was imported. Review the file again.", 409, "zone_import_conflict")
            for row in rows:
                zones.insert(db, current, ZoneFields(name=row["name"], status=row["status"]))
            db.commit()
            return {"imported": len(rows)}
        db.commit()
        return dict(rows=rows, valid=valid, digest=fingerprint)
    return zones.transaction(db, work)


def export(db, actor, query, status, format):
    def work():
        zones.authorize(db, actor)
        records = list(db.scalars(select(Zone).where(*zones.predicates(query, status))
                                  .order_by(Zone.created_at.desc(), Zone.id.desc()).limit(EXPORT_LIMIT + 1)))
        if len(records) > EXPORT_LIMIT:
            raise zones.ZoneError("Export exceeds 1,000 matching records. Narrow the name/status filters.", 422, "zone_export_limit")
        rows = [HEADERS]
        for record in records:
            projected = zones.projection(db, record)
            rows.append([safe_text(record.name), record.status, safe_text(projected["createdBy"]),
                         record.created_at.isoformat(), safe_text(projected["updatedBy"]), record.updated_at.isoformat()])
        db.commit()
        return rows
    rows = zones.transaction(db, work)
    if format == "csv":
        output = io.StringIO()
        csv.writer(output).writerows(rows)
        return output.getvalue().encode("utf-8-sig")
    book = Workbook()
    for row in rows:
        book.active.append(row)
        for cell in book.active[book.active.max_row]:
            cell.data_type = "s"
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
