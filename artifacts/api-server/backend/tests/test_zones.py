"""Zone API and transfer regressions: synthetic PostgreSQL fixture only."""
import io
import uuid
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.db.models import AuditEvent
from app.db.zone_models import Zone
from app.services.zone_transfer import parse, MAX_BYTES
from app.services.zones import ZoneError
from test_sessions import client, create_user, login
from test_reporting import admin_headers

BASE = "/api/v1/admin/zones"


def add(api, headers, name="North", status="active"):
    result = api.post(BASE, headers=headers, json={"name": name, "status": status})
    assert result.status_code == 201, result.text
    return result.json()


def workbook(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def replace_zip(data, name, transform):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, "w") as target:
        for item in source.infolist():
            content = source.read(item)
            target.writestr(item.filename, transform(content) if item.filename == name else content)
    return output.getvalue()


def review(api, headers, data, filename="zones.csv"):
    return api.post(BASE + "/import/review", headers={**headers, "Content-Type": "application/octet-stream"},
                    params={"filename": filename}, content=data)


def commit(api, headers, data, digest, filename="zones.csv"):
    return api.post(BASE + "/import/commit", headers=headers,
                    params={"filename": filename, "digest": digest, "confirm": "true"}, content=data)


def test_crud_soft_delete_creator_version_and_filters(client):
    api, db, _ = client
    headers, actor = admin_headers(api, db)
    row = add(api, headers, "  North  ")
    assert row["name"] == "North" and row["createdBy"] == "Super Admin"
    assert row["createdAt"].endswith("Z") or "+00:00" in row["createdAt"]
    assert api.post(BASE, headers=headers, json={"name": "NORTH", "status": "inactive"}).status_code == 409
    edited = api.post(f"{BASE}/{row['id']}/edit", headers=headers,
                     json={"name": "North revised", "status": "active", "expected_version": 1}).json()
    assert edited["version"] == 2 and edited["createdAt"] == row["createdAt"]
    assert api.post(f"{BASE}/{row['id']}/status", headers=headers,
                    json={"status": "inactive", "expected_version": 1}).status_code == 409
    inactive = api.post(f"{BASE}/{row['id']}/status", headers=headers,
                        json={"status": "inactive", "expected_version": 2})
    assert inactive.status_code == 200
    for name in ("South", "East", "West"):
        add(api, headers, name)
    page = api.get(BASE, headers=headers, params={"query": "north", "status": "inactive", "limit": 1}).json()
    assert page["total"] == 4 and page["filtered"] == 1 and page["items"][0]["id"] == row["id"]
    assert api.get(BASE, headers=headers, params={"query": "%"}).json()["filtered"] == 0
    assert api.get(BASE, headers=headers, params={"status": "active", "limit": 1, "offset": 1}).json()["filtered"] == 3
    assert api.post(f"{BASE}/{row['id']}/delete", headers=headers, json={"expected_version": 2}).status_code == 409
    deleted = api.post(f"{BASE}/{row['id']}/delete", headers=headers, json={"expected_version": 3})
    assert deleted.status_code == 200 and deleted.json()["version"] == 4
    db.expire_all()
    stored = db.get(Zone, uuid.UUID(row["id"]))
    assert stored.deleted_by == actor.id and stored.deleted_at is not None
    assert stored.updated_at == stored.deleted_at and stored.updated_by == actor.id and stored.created_by == actor.id
    assert api.get(f"{BASE}/{row['id']}", headers=headers).status_code == 404
    assert api.get(BASE, headers=headers).json()["total"] == 3
    assert "North revised" not in api.get(BASE + "/export", headers=headers).text
    replacement = add(api, headers, "North revised")
    assert replacement["id"] != row["id"]
    event = db.scalar(select(AuditEvent).where(AuditEvent.action == "zone_delete"))
    assert event.actor_id == actor.id and event.resource_id == stored.id and event.session_id


def test_authorization_forgery_expired_session_and_validation(client):
    api, db, _ = client
    assert api.get(BASE).status_code == 401
    mr = create_user(db, "zone-mr@example.com")
    token = login(api, mr.email).json()["access_token"]
    headers = {"Authorization": "Bearer " + token}
    for route in (BASE, BASE + "/export"):
        assert api.get(route, headers=headers).status_code == 403
    assert review(api, headers, b"Zone Name,Status\nSecret,Active").status_code == 403
    headers, _ = admin_headers(api, db)
    for extra in ({"createdBy": "Forged"}, {"created_by": str(uuid.uuid4())}, {"deleted_at": "now"}, {"version": 99}):
        assert api.post(BASE, headers=headers, json={"name": "Forgery", "status": "active", **extra}).status_code == 422
    for name in ("", " " * 3, "x" * 201):
        assert api.post(BASE, headers=headers, json={"name": name, "status": "active"}).status_code == 422
    assert db.scalar(select(func.count()).select_from(Zone)) == 0
    api.post("/api/v1/auth/logout", headers={"Origin": "http://testserver"}, json={})
    assert api.post(BASE, headers=headers, json={"name": "Expired", "status": "active"}).status_code == 401


def test_trash_restore_metadata_versions_and_privacy(client):
    api, db, _ = client
    headers, actor = admin_headers(api, db)
    row = add(api, headers, "Restore original", "inactive")
    zone_id = row["id"]
    deleted = api.post(f"{BASE}/{zone_id}/delete", headers=headers, json={"expected_version": 1}).json()
    add(api, headers, "Current")
    trash = api.get(BASE + "/trash", headers=headers, params={"query": "original", "status": "inactive", "limit": 1})
    assert trash.status_code == 200 and trash.headers["cache-control"] == "no-store"
    page = trash.json()
    assert page["total"] == page["filtered"] == 1
    assert page["items"][0] == {**deleted, "deletedBy": "Super Admin", "deletedAt": deleted["updatedAt"]}
    assert api.get(BASE + "/trash", headers=headers, params={"query": "%"}).json()["filtered"] == 0
    assert api.get(BASE + "/trash", headers=headers, params={"status": "active"}).json()["filtered"] == 0
    assert api.get(BASE + "/trash", headers=headers, params={"offset": 1}).json()["items"] == []
    assert "Restore original" not in api.get(BASE + "/export", headers=headers).text
    assert api.get(f"{BASE}/{zone_id}", headers=headers).status_code == 404
    assert api.post(f"{BASE}/{zone_id}/restore", headers=headers, json={"expected_version": 1}).json()["error"]["code"] == "zone_stale"
    for extra in ({"createdBy": "Forged"}, {"status": "active"}, {"deletedBy": "Forged"}):
        assert api.post(f"{BASE}/{zone_id}/restore", headers=headers, json={"expected_version": 2, **extra}).status_code == 422
    assert api.post(f"{BASE}/{zone_id}/restore", headers=headers, json={"expected_version": "2"}).status_code == 422
    result = api.post(f"{BASE}/{zone_id}/restore", headers=headers, json={"expected_version": 2})
    assert result.status_code == 200 and result.headers["cache-control"] == "no-store"
    restored = result.json()
    assert restored["version"] == 3 and restored["status"] == "inactive"
    assert restored["createdAt"] == row["createdAt"] and restored["createdBy"] == row["createdBy"]
    assert restored["updatedAt"] > deleted["updatedAt"]
    db.expire_all()
    stored = db.get(Zone, uuid.UUID(zone_id))
    assert stored.deleted_at is None and stored.deleted_by is None
    assert stored.created_by == actor.id and stored.updated_by == actor.id
    assert api.get(BASE + "/trash", headers=headers).json()["total"] == 0
    assert api.get(f"{BASE}/{zone_id}", headers=headers).json() == restored
    assert "Restore original" in api.get(BASE + "/export", headers=headers).text
    assert api.post(f"{BASE}/{zone_id}/restore", headers=headers, json={"expected_version": 3}).status_code == 409
    assert api.post(f"{BASE}/{uuid.uuid4()}/restore", headers=headers, json={"expected_version": 2}).status_code == 404
    events = db.scalars(select(AuditEvent).where(AuditEvent.resource_id == stored.id, AuditEvent.action == "zone_restore")).all()
    assert len(events) == 1 and events[0].actor_id == actor.id and events[0].session_id
    assert db.scalar(select(AuditEvent.id).where(AuditEvent.resource_id == stored.id, AuditEvent.action == "zone_delete"))


@pytest.mark.parametrize("status", ["active", "inactive"])
def test_restore_duplicate_name_rolls_back_zone_and_audit(client, status):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    row = add(api, headers, "Reused")
    deleted = api.post(f"{BASE}/{row['id']}/delete", headers=headers, json={"expected_version": 1}).json()
    replacement = add(api, headers, "REUSED", status)
    result = api.post(f"{BASE}/{row['id']}/restore", headers=headers, json={"expected_version": 2})
    assert result.status_code == 409 and result.json()["error"]["code"] == "zone_duplicate"
    tombstone = api.get(BASE + "/trash", headers=headers).json()["items"][0]
    assert tombstone["version"] == 2 and tombstone["updatedAt"] == deleted["updatedAt"]
    assert tombstone["deletedAt"] == deleted["updatedAt"]
    assert not db.scalar(select(AuditEvent.id).where(AuditEvent.action == "zone_restore"))
    # Conflict must be resolved explicitly; restore never renames either record.
    assert api.post(f"{BASE}/{replacement['id']}/edit", headers=headers,
                    json={"name": "Freed", "status": status, "expected_version": 1}).status_code == 200
    assert api.post(f"{BASE}/{row['id']}/restore", headers=headers, json={"expected_version": 2}).status_code == 200


def test_restore_audit_failure_is_atomic(client, monkeypatch):
    from sqlalchemy.exc import SQLAlchemyError
    from app.services import zones
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    row = add(api, headers)
    api.post(f"{BASE}/{row['id']}/delete", headers=headers, json={"expected_version": 1})
    original = zones.audit
    def fail_audit(db, actor, row, action):
        original(db, actor, row, action)
        db.flush()
        raise SQLAlchemyError("Synthetic audit failure")
    monkeypatch.setattr(zones, "audit", fail_audit)
    assert api.post(f"{BASE}/{row['id']}/restore", headers=headers, json={"expected_version": 2}).status_code == 503
    db.expire_all()
    stored = db.get(Zone, uuid.UUID(row["id"]))
    assert stored.deleted_at and stored.deleted_by and stored.version == 2
    assert not db.scalar(select(AuditEvent.id).where(AuditEvent.action == "zone_restore"))


def test_trash_restore_auth_and_live_identity_revalidation(client):
    from app.db.models import AuthSession
    from app.services.auth import AuthError, Identity
    from app.services import zones
    from app.schemas.zones import ZoneVersion
    api, db, _ = client
    path = f"{BASE}/{uuid.uuid4()}/restore"
    assert api.get(BASE + "/trash").status_code == 401
    assert api.post(path, json={"expected_version": 1}).status_code == 401
    mr = create_user(db, "trash-mr@example.com")
    headers = {"Authorization": "Bearer " + login(api, mr.email).json()["access_token"]}
    assert api.get(BASE + "/trash", headers=headers).status_code == 403
    assert api.post(path, headers=headers, json={"expected_version": 1}).status_code == 403
    headers, actor = admin_headers(api, db)
    row = add(api, headers)
    api.post(f"{BASE}/{row['id']}/delete", headers=headers, json={"expected_version": 1})
    session = db.scalar(select(AuthSession).where(AuthSession.user_id == actor.id, AuthSession.status == "ACTIVE"))
    snapshot = Identity(actor, session_id=session.id)
    # Sensitive services must not trust an identity verified before revocation.
    from app.core.security import utcnow
    session.status, session.revoked_at = "REVOKED", utcnow()
    db.commit()
    with pytest.raises(AuthError):
        zones.restore(db, snapshot, uuid.UUID(row["id"]), ZoneVersion(expected_version=2))
    with pytest.raises(AuthError):
        zones.listing(db, snapshot, "", "all", 10, 0, deleted=True)
    assert api.get(BASE + "/trash", headers=headers).status_code == 401
    assert api.post(f"{BASE}/{row['id']}/restore", headers=headers, json={"expected_version": 2}).status_code == 401
    assert db.get(Zone, uuid.UUID(row["id"])).deleted_at is not None


def test_review_create_only_commit_rechecks_and_atomic_rollback(client):
    api, db, _ = client
    headers, actor = admin_headers(api, db)
    data = b"\xef\xbb\xbfZone Name,Status,Created By,Created At,Updated By,Updated At\r\nImported,Active,Forged,1900,Forged,1900\r\nSecond,Inactive,Forged,1900,Forged,1900"
    result = review(api, headers, data)
    assert result.status_code == 200, result.text
    report = result.json()
    assert report["valid"] and db.scalar(select(func.count()).select_from(Zone)) == 0
    assert commit(api, headers, data + b"\n", report["digest"]).status_code == 409
    add(api, headers, "Second")
    assert commit(api, headers, data, report["digest"]).status_code == 409
    assert db.scalar(select(func.count()).select_from(Zone)) == 1
    assert not db.scalar(select(Zone).where(Zone.name == "Imported"))
    duplicate = review(api, headers, b"Zone Name,Status\nDUP,active\ndup,active").json()
    assert not duplicate["valid"] and duplicate["rows"][1]["errors"]
    clean = b"Zone Name,Status\nImported,active"
    report = review(api, headers, clean).json()
    assert commit(api, headers, clean, report["digest"]).json() == {"imported": 1}
    stored = db.scalar(select(Zone).where(Zone.name == "Imported"))
    assert stored.created_by == actor.id and stored.updated_by == actor.id


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_export_all_filtered_matches_and_safe_roundtrip(client, format):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    for name in ("=SUM(1)", "'=literal", "+north", "Plain"):
        add(api, headers, name)
    add(api, headers, "Inactive", "inactive")
    listing = api.get(BASE, headers=headers, params={"limit": 2, "status": "active"}).json()
    assert len(listing["items"]) == 2 and listing["filtered"] == 4
    response = api.get(BASE + "/export", headers=headers, params={"format": format, "status": "active"})
    assert response.status_code == 200 and response.headers["content-disposition"].endswith(f'.{format}"')
    if format == "xlsx":
        assert "spreadsheetml.sheet" in response.headers["content-type"]
        book = load_workbook(io.BytesIO(response.content))
        assert all(cell.data_type != "f" for row in book.active for cell in row)
        book.close()
    else:
        assert response.headers["content-type"].startswith("text/csv")
    rows = parse(response.content, "zones." + format)
    assert {row["name"] for row in rows} == {"=SUM(1)", "'=literal", "+north", "Plain"}
    for row in list(db.scalars(select(Zone))):
        api.post(f"{BASE}/{row.id}/delete", headers=headers, json={"expected_version": row.version})
    report = review(api, headers, response.content, "zones." + format).json()
    assert report["valid"]
    assert commit(api, headers, response.content, report["digest"], "zones." + format).json()["imported"] == 4


def test_xlsx_bom_invalid_and_limits(client):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    valid = workbook([["Zone Name", "Status"], ["Workbook Zone", "Active"]])
    assert review(api, headers, valid, "zones.xlsx").json()["valid"]
    for data, name in ((b"\xff", "zones.csv"), (b"Not a zip", "zones.xlsx"), (valid, "zones.xls"),
                       (b"Wrong,Status\nName,active", "zones.csv"),
                       (b'Zone Name,Status\n"unterminated,active', "zones.csv"),
                       (workbook([["Zone Name", "Status"], ["=1+1", "Active"]]), "zones.xlsx")):
        assert review(api, headers, data, name).status_code == 422
    assert review(api, headers, b"x" * (MAX_BYTES + 1)).status_code == 413
    assert review(api, headers, b"Zone Name,Status\n" + b"Zone,Active\n" * 1001).status_code == 422
    hostile = io.BytesIO()
    with zipfile.ZipFile(hostile, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/sharedStrings.xml", "x" * (9 * 1024 * 1024))
    assert review(api, headers, hostile.getvalue(), "bomb.xlsx").status_code == 422


def test_export_limit_and_both_size_layers(client):
    api, db, _ = client
    headers, actor = admin_headers(api, db)
    db.add_all([Zone(name=f"Bound {index}", status="active", created_by=actor.id, updated_by=actor.id) for index in range(1001)])
    db.commit()
    assert api.get(BASE + "/export", headers=headers).json()["error"]["code"] == "zone_export_limit"
    assert api.get(BASE + "/export", headers=headers, params={"query": "Bound 1000"}).status_code == 200
    assert review(api, headers, b"Zone Name,Status\nLarge,active",
                  "zones.csv").status_code == 200
    # Chunked upload uses actual bytes, not just the declared Content-Length.
    response = api.post(BASE + "/import/review", params={"filename": "zones.csv"}, headers=headers,
                        content=iter([b"x" * 1024 * 1024] * 3))
    assert response.status_code == 413
    assert api.post(BASE, headers=headers, content=b"x" * (1024 * 1024 + 1)).status_code == 413


def test_workbook_rejects_macros_links_entities_coordinates_and_does_not_trust_dimensions():
    valid = workbook([["Zone Name", "Status"], ["Visible", "active"], ["Hidden", "inactive"]])
    understated = replace_zip(valid, "xl/worksheets/sheet1.xml",
                               lambda data: data.replace(b'ref="A1:B3"', b'ref="A1:B2"'))
    assert len(parse(understated, "zones.xlsx")) == 2
    cases = [
        replace_zip(valid, "xl/worksheets/sheet1.xml", lambda data: data.replace(b'r="A3"', b'r="A999999"')),
        replace_zip(valid, "xl/worksheets/sheet1.xml", lambda data: data.replace(b'r="A3"', b'r="ZZ3"')),
        replace_zip(valid, "xl/worksheets/sheet1.xml", lambda data: b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]>' + data),
        replace_zip(valid, "xl/_rels/workbook.xml.rels", lambda data: data.replace(b'Target=', b'TargetMode="External" Target=', 1)),
        replace_zip(valid, "[Content_Types].xml", lambda data: data.replace(b'spreadsheetml.sheet.main+xml', b'spreadsheetml.macroEnabled.main+xml')),
        workbook([["Zone Name", "Status"], ["=" + "1", "active"]]),
        workbook([["Zone Name", "Status"], ["x" * 10001, "active"]]),
    ]
    for data in cases:
        with pytest.raises(ZoneError):
            parse(data, "zones.xlsx")
