"""Directory tombstones preserve ownership, files, identifiers and historical references."""
import uuid
import csv
import io
import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from app.db.models import Patient, AuditEvent
from app.db.file_models import FileRecord
from app.db.doctor_models import DoctorDirectory
from app.db.patient_models import PatientDirectory
from test_sessions import client
from test_patients import setup_patient, fields, add, edit, csv_file, review, commit, BASE
from test_doctors import fields as doctor_fields, add as add_doctor, BASE as DOCTORS
from test_mrs import add as add_mr, fields as mr_fields


def delete(api, headers, base, record, **extra):
    return api.post(f"{base}/{record['id']}/delete", headers=headers,
                    json={"expected_version": record["version"], **extra})


def export(api, headers, base, format, **filters):
    return api.get(base + "/export", headers={**headers, "X-Download-Initiation": str(uuid.uuid4())},
                   params={"format": format, **filters})


def assert_empty_export(response, format):
    assert response.status_code == 200, response.text[:100]
    rows = (list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")))) if format == "csv" else
            list(load_workbook(io.BytesIO(response.content), read_only=True).active.values))
    assert len(rows) == 1, "Only the unchanged export header should remain."


def test_doctor_delete_historical_references_and_import_race(client):
    api, db, _ = client
    headers, actor, mr, zone, doctor = setup_patient(api, db)
    patient = add(api, headers, fields(doctor))
    owner = db.get(Patient, uuid.UUID(patient["id"]))
    before = (owner.assigned_mr_id, owner.is_active, owner.version)
    balances = "/api/v1/admin/opening-balances"
    balance = api.post(balances, headers=headers, json=dict(
        doctorId=doctor["id"], startYear=2025, endYear=2026, amount="12.34", status="active")).json()
    data = csv_file([fields(doctor, name="Reviewed Patient")], doctor)
    report = review(api, headers, data).json()
    assert report["valid"]
    assert delete(api, headers, DOCTORS, doctor, deleted_by=str(uuid.uuid4())).status_code == 422
    assert delete(api, {}, DOCTORS, doctor).status_code == 401
    assert delete(api, headers, DOCTORS, {**doctor, "version": 9}).status_code == 409
    response = delete(api, headers, DOCTORS, doctor)
    assert response.status_code == 200 and response.json()["version"] == doctor["version"] + 1
    assert response.json()["status"] == doctor["status"]
    row = db.get(DoctorDirectory, uuid.UUID(doctor["id"]))
    assert row.deleted_at.tzinfo and row.deleted_by == actor.id and row.updated_by == actor.id
    assert delete(api, headers, DOCTORS, doctor).status_code == 404
    for suffix in ("", "/edit", "/status", "/contact"):
        if not suffix:
            assert api.get(DOCTORS + "/" + doctor["id"], headers=headers).status_code == 404
        else:
            body = {**doctor_fields(mr), "expected_version": 2} if suffix == "/edit" else (
                {"expected_version": 2, "status": "inactive"} if suffix == "/status" else
                {"expected_version": 2, "contactRequirement": "optional"})
            assert api.post(DOCTORS + "/" + doctor["id"] + suffix, headers=headers, json=body).status_code == 404
    for status in ("all", "active", "inactive"):
        page = api.get(DOCTORS, headers=headers, params={"status": status}).json()
        assert page["total"] == page["filtered"] == 0 and not page["items"]
    assert api.get(DOCTORS + "/filters", headers=headers).json()["states"] == []
    assert api.post(DOCTORS + "/bulk", headers=headers, json=dict(
        operation="verification", verification="verified", selected=[dict(id=doctor["id"], expected_version=2)])).status_code == 404
    for path, params in ((BASE + "/references", {"include_saved": doctor["id"]}),
                         (balances + "/references", {"balance_id": balance["id"]}),
                         ("/api/v1/admin/mrs/" + mr["id"] + "/doctors", {})):
        result = api.get(path, headers=headers, params=params)
        assert result.status_code == 200, result.text
        assert result.json()["total"] == 0 and not result.json()["items"]
    current = api.get(BASE + "/" + patient["id"], headers=headers).json()
    assert current["doctorId"] == doctor["id"] and current["doctorName"] == doctor["name"]
    assert any("deleted" in w for w in current["assignmentWarnings"])
    # A Doctor tombstone is not a Patient tombstone. Retained MR/Zone filters
    # must still find and export the live Patient without making a live choice.
    for filters in ({"mr_id": mr["id"]}, {"zone_id": zone},
                    {"mr_id": mr["id"], "zone_id": zone, "query": patient["name"], "status": "active"}):
        retained_page = api.get(BASE, headers=headers, params={**filters, "limit": 1, "offset": 0}).json()
        assert retained_page["total"] == retained_page["filtered"] == 1
        assert retained_page["items"][0]["id"] == patient["id"]
        empty_page = api.get(BASE, headers=headers, params={**filters, "limit": 1, "offset": 1}).json()
        assert empty_page["filtered"] == 1 and not empty_page["items"]
        for format in ("csv", "xlsx"):
            result = export(api, headers, BASE, format, **filters)
            assert result.status_code == 200
            exported = (list(csv.reader(io.StringIO(result.content.decode("utf-8-sig")))) if format == "csv" else
                        list(load_workbook(io.BytesIO(result.content), read_only=True).active.values))
            assert len(exported) == 2 and patient["code"] in exported[1]
        assert api.get(DOCTORS, headers=headers, params=filters).json()["filtered"] == 0
    assert edit(api, headers, current, fields(doctor, name="Retained Patient")).status_code == 200
    owner = db.get(Patient, uuid.UUID(patient["id"]))
    assert (owner.assigned_mr_id, owner.is_active) == before[:2]
    retained = api.get(balances + "/" + balance["id"], headers=headers).json()
    assert retained["doctorName"] == doctor["name"] and not retained["doctorUsable"]
    assert api.post(balances + "/" + balance["id"] + "/edit", headers=headers,
                    json={**{k: balance[k] for k in ("doctorId", "startYear", "endYear", "amount", "status")},
                          "expected_version": balance["version"]}).status_code == 200
    assert api.post(BASE, headers=headers, json=fields(doctor, name="New Patient")).status_code == 409
    assert commit(api, headers, data, report["digest"]).status_code == 409
    assert not review(api, headers, data).json()["valid"]
    assert api.post(DOCTORS, headers=headers, json=doctor_fields(mr)).status_code == 409
    for format in ("csv", "xlsx"):
        exported = export(api, headers, DOCTORS, format)
        assert_empty_export(exported, format)
    event = db.scalar(select(AuditEvent).where(AuditEvent.action == "doctor_directory_delete"))
    assert event.actor_id == actor.id and str(event.resource_id) == doctor["id"]


def test_patient_delete_preserves_owner_files_and_ignores_later_shift(client):
    api, db, _ = client
    headers, actor, mr, zone, doctor = setup_patient(api, db)
    patient = add(api, headers, fields(doctor))
    owner = db.get(Patient, uuid.UUID(patient["id"]))
    before = (owner.assigned_mr_id, owner.is_active, owner.version)
    fid = uuid.uuid4()
    key = f"patients/{owner.id}/documents/{fid}.pdf"
    db.add(FileRecord(id=fid, patient_id=owner.id, category="documents", object_key=key,
                      display_name="Synthetic retained file", content_type="application/pdf",
                      uploader_id=actor.id, state="deleted"))
    db.commit()
    assert delete(api, headers, BASE, patient, deleted_at="2020-01-01").status_code == 422
    assert delete(api, headers, BASE, {**patient, "version": 9}).status_code == 409
    assert delete(api, headers, BASE, patient).status_code == 200
    owner = db.get(Patient, owner.id)
    assert (owner.assigned_mr_id, owner.is_active, owner.version) == (*before[:2], before[2] + 1)
    row = db.get(PatientDirectory, owner.id)
    snapshot = (row.doctorId, row.status, row.updated_at, owner.version, row.updated_by)
    assert row.deleted_at.tzinfo and row.deleted_by == actor.id
    file = db.get(FileRecord, fid)
    assert file.object_key == key and file.state == "deleted" and file.patient_id == owner.id
    assert delete(api, headers, BASE, patient).status_code == 404
    assert api.get(BASE + "/" + patient["id"], headers=headers).status_code == 404
    assert edit(api, headers, {**patient, "version": 2}, fields(doctor)).status_code == 404
    assert api.post(BASE + "/" + patient["id"] + "/status", headers=headers,
                    json=dict(expected_version=2, status="inactive")).status_code == 404
    assert api.get(BASE, headers=headers).json()["total"] == 0
    assert api.post(BASE, headers=headers, json=fields(doctor)).status_code == 409
    for format in ("csv", "xlsx"):
        exported = export(api, headers, BASE, format)
        assert_empty_export(exported, format)
    other = add_mr(api, headers, mr_fields(mr["hq"], zone, name="Another MR",
                                         username="another.mr", code="ANOTHER"))["record"]
    response = api.post(DOCTORS + "/" + doctor["id"] + "/edit", headers=headers,
                        json={**doctor_fields(mr), "mrId": other["id"], "expected_version": doctor["version"]})
    assert response.status_code == 200, response.text
    db.refresh(row); db.refresh(owner)
    assert (row.doctorId, row.status, row.updated_at, owner.version, row.updated_by) == snapshot
    assert owner.assigned_mr_id == before[0]


def test_opening_balance_import_rejects_deleted_doctor_at_commit(client):
    api, db, _ = client
    headers, _, _, _, doctor = setup_patient(api, db)
    from test_opening_balances import file, upload, BASE as balances
    data = file(doctor)
    report = upload(api, headers, data).json()
    assert report["valid"]
    assert delete(api, headers, DOCTORS, doctor).status_code == 200
    assert upload(api, headers, data, digest=report["digest"]).status_code == 409
    assert not upload(api, headers, data).json()["valid"]
    assert api.get(balances, headers=headers).json()["total"] == 0


def test_deleted_patient_code_and_duplicate_remain_reserved(client):
    api, db, _ = client
    headers, _, _, _, doctor = setup_patient(api, db)
    patient = add(api, headers, fields(doctor))
    assert delete(api, headers, BASE, patient).status_code == 200
    # Portable import with different business identity cannot reuse the saved code.
    data = csv_file([fields(doctor, name="Different identity")], doctor)
    data = data.replace(b"PAT-LEGACY-000000", patient["code"].encode())
    assert not review(api, headers, data).json()["valid"]


def test_live_inactive_records_and_paged_reference_counts_survive(client):
    api, db, _ = client
    headers, _, mr, _, doctor = setup_patient(api, db)
    inactive = add_doctor(api, headers, doctor_fields(mr, name="Inactive retained",
                                                     registrationNumber="REG-INACTIVE", status="inactive"))
    live = add_doctor(api, headers, doctor_fields(mr, name="Live remaining", registrationNumber="REG-LIVE"))
    assert delete(api, headers, DOCTORS, doctor).status_code == 200
    page = api.get(DOCTORS, headers=headers, params=dict(status="inactive")).json()
    assert page["total"] == 2 and page["filtered"] == 1 and page["items"][0]["id"] == inactive["id"]
    choices = api.get(BASE + "/references", headers=headers, params=dict(limit=1, offset=0, include_saved=doctor["id"])).json()
    assert choices["total"] == 2 and len(choices["items"]) == 1 and choices["items"][0]["id"] == inactive["id"]
    second = api.get(BASE + "/references", headers=headers, params=dict(limit=1, offset=1)).json()
    assert second["total"] == 2 and len(second["items"]) == 1 and second["items"][0]["id"] == live["id"]


@pytest.mark.parametrize("resource", ("doctor", "patient"))
def test_delete_audit_failure_rolls_back_all_metadata(client, monkeypatch, resource):
    api, db, _ = client
    headers, _, _, _, doctor = setup_patient(api, db)
    patient = add(api, headers, fields(doctor))
    from app.services import doctors, patients
    service = doctors if resource == "doctor" else patients
    record, base, model = (doctor, DOCTORS, DoctorDirectory) if resource == "doctor" else (patient, BASE, PatientDirectory)
    def fail(*args):
        raise service.DoctorError() if resource == "doctor" else service.PatientError()
    monkeypatch.setattr(service, "audit", fail)
    assert delete(api, headers, base, record).status_code == 503
    row = db.get(model, uuid.UUID(record["id"]))
    assert row.deleted_at is None and row.deleted_by is None
    assert (row.version if resource == "doctor" else db.get(Patient, row.id).version) == record["version"]
