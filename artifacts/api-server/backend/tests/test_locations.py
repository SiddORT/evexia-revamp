"""Synthetic isolated PostgreSQL only; never uses managed data or credentials."""
import csv
from collections import OrderedDict
import io
import uuid
import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from app.db.models import AuditEvent
from app.db.location_models import StorageLocation
from app.services.location_transfer import parse, MAX_BYTES
from app.services.locations import LocationError
from test_sessions import client, create_user, login
from test_reporting import admin_headers
from test_zones import workbook, replace_zip

BASE = "/api/v1/admin/storage-locations"
LEGACY = ["Storage Location", "Address", "Status"]
AUDIT = ["Created By", "Created At", "Updated By", "Updated At"]
STRUCTURED = LEGACY + ["PIN Code", "Country Code", "State/UT", "City"]
EMPTY_GEO = dict(pincode="", countryCode="", state="", city="")
PUNE = dict(pincode="411001", countryCode="IN", state="Maharashtra", city="Pune")


def transfer_file(rows, format):
    if format == "xlsx":
        return workbook(rows)
    output = io.StringIO()
    csv.writer(output).writerows(rows)
    return output.getvalue().encode("utf-8-sig")


@pytest.mark.parametrize("format", ["csv", "xlsx"])
@pytest.mark.parametrize("structured,audited", [(False, False), (False, True), (True, False), (True, True)])
def test_all_exact_backup_shapes_preserve_explicit_geography(client, format, structured, audited):
    api, db, _ = client
    headers, actor = admin_headers(api, db)
    columns = (STRUCTURED if structured else LEGACY) + (AUDIT if audited else [])
    # Legacy address text is deliberately recognizable: it must never be parsed
    # into authoritative geography, including on seven-column audited backups.
    cells = ["Backup location", "Pune Maharashtra India 411001", "Inactive"]
    cells += ["411001", "IN", "  maharashtra ", " Pune "] if structured else []
    cells += ["Forged", "1900", "Forged", "1900"] if audited else []
    data = transfer_file([columns, cells], format)
    filename = "locations." + format
    response = review(api, headers, data, filename)
    assert response.status_code == 200, response.text
    report = response.json()
    expected = PUNE if structured else EMPTY_GEO
    assert report["valid"] and len(report["rows"]) == 1
    assert {key: report["rows"][0][key] for key in EMPTY_GEO} == expected
    assert not report["rows"][0]["errors"]
    assert db.scalar(select(func.count()).select_from(StorageLocation)) == 0
    assert commit(api, headers, data, report["digest"], filename).json() == {"imported": 1}
    stored = db.scalar(select(StorageLocation))
    assert {key: getattr(stored, key) for key in EMPTY_GEO} == expected
    assert stored.address == cells[1] and stored.status == "inactive"
    assert stored.created_by == stored.updated_by == actor.id and stored.created_at.year > 1900
    response = api.get(f"{BASE}/{stored.id}", headers=headers).json()
    assert {key: response[key] for key in EMPTY_GEO} == expected
    assert response["stateCode"] == ("27" if structured else None)


@pytest.mark.parametrize("format", ["csv", "xlsx"])
@pytest.mark.parametrize("columns", [LEGACY, LEGACY + AUDIT, STRUCTURED, STRUCTURED + AUDIT])
def test_header_shapes_are_exact_not_just_matching_width(format, columns):
    mutations = [columns[::-1], [columns[0] + " ", *columns[1:]],
                 [*columns[:-1], columns[-1].lower()], [*columns, "ID"]]
    for changed in mutations:
        with pytest.raises(LocationError) as exc:
            parse(transfer_file([changed, ["x"] * len(changed)], format), "locations." + format)
        assert exc.value.code == "location_invalid_file"


@pytest.mark.parametrize("alias,canonical,code", [
    ("27", "Maharashtra", "27"), ("maha-rashtra", "Maharashtra", "27"),
    ("Orissa", "Odisha", "21"), ("Uttaranchal", "Uttarakhand", "05"),
    ("Pondicherry", "Puducherry", "34"), ("NCT of Delhi", "Delhi", "07"),
    ("Andaman & Nicobar Islands", "Andaman and Nicobar Islands", "35"),
    ("Dadra and Nagar Haveli", "Dadra and Nagar Haveli and Daman and Diu", "26"),
    ("Daman and Diu", "Dadra and Nagar Haveli and Daman and Diu", "26"),
])
def test_state_aliases_are_canonical_in_review_and_saved_response(client, alias, canonical, code):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    data = transfer_file([STRUCTURED, ["Alias", "Office", "active", "411001", "IN", alias, "City"]], "csv")
    report = review(api, headers, data).json()
    assert report["valid"] and report["rows"][0]["state"] == canonical
    assert commit(api, headers, data, report["digest"]).status_code == 200
    saved = api.get(BASE, headers=headers).json()["items"][0]
    assert saved["state"] == canonical and saved["stateCode"] == code


@pytest.mark.parametrize("field,value", [
    ("pincode", "x" * 21), ("state", "x" * 101), ("city", "x" * 101),
    ("countryCode", "USA"), ("countryCode", "ZZ"), ("countryCode", "in"),
    ("countryCode", ""), ("state", "Nul\x00state"), ("city", "Nul\x00city"),
    ("pincode", "Nul\x00pin"),
])
def test_geography_bounds_reject_api_and_import_without_partial_writes(client, field, value):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    geo = dict(pincode="12345", countryCode="US", state="Region", city="City")
    geo[field] = value
    body = dict(name="Invalid geography", address="Office", status="active", **geo)
    assert api.post(BASE, headers=headers, json=body).status_code == 422
    data = transfer_file([STRUCTURED, ["Valid first", "Office", "active", *PUNE.values()],
                          [body[key] for key in ("name", "address", "status", *EMPTY_GEO)]], "csv")
    report = review(api, headers, data).json()
    assert not report["valid"] and not report["rows"][0]["errors"] and report["rows"][1]["errors"]
    assert commit(api, headers, data, report["digest"]).status_code == 409
    assert db.scalar(select(func.count()).select_from(StorageLocation)) == 0
    assert db.scalar(select(func.count()).select_from(AuditEvent).where(AuditEvent.action == "location_create")) == 0


@pytest.mark.parametrize("pin", ["011001", "12345", "1234567", "１２３４５６", "12345a"])
def test_indian_pin_validation_is_not_generic_postal_text(client, pin):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    assert api.post(BASE, headers=headers, json=dict(name="Bad PIN", address="Office", status="active",
                                                    **{**PUNE, "pincode": pin})).status_code == 422
    row = parse(transfer_file([STRUCTURED, ["Bad PIN", "Office", "active", pin, "IN", "Maharashtra", "Pune"]], "csv"),
                "locations.csv")[0]
    assert row["errors"]


def test_maximum_foreign_geography_and_optional_fields_are_accepted(client):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    geo = dict(pincode="x" * 20, countryCode="US", state="s" * 100, city="c" * 100)
    response = api.post(BASE, headers=headers, json=dict(name="Bounded", address="Office", status="active", **geo))
    assert response.status_code == 201, response.text
    assert {key: response.json()[key] for key in EMPTY_GEO} == geo
    assert response.json()["stateCode"] is None
    for partial in ({}, {"countryCode": "IN"}, {"countryCode": "IN", "state": "Maharashtra"}):
        row = api.post(BASE, headers=headers, json=dict(name="Optional " + str(len(partial)), address="Office",
                                                       status="active", **partial))
        assert row.status_code == 201, row.text


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_current_formula_safe_geography_roundtrip_and_digest_binding(client, format):
    api, db, _ = client
    headers, actor = admin_headers(api, db)
    geo = dict(pincode="+12345", countryCode="US", state="'=literal", city="@City")
    body = dict(name="=SUM(1)", address="-Office", status="active", **geo)
    saved = api.post(BASE, headers=headers, json=body)
    assert saved.status_code == 201, saved.text
    exported = api.get(BASE + "/export", headers=headers, params={"format": format}).content
    if format == "xlsx":
        book = load_workbook(io.BytesIO(exported))
        assert list(next(book.active.values)) == STRUCTURED + AUDIT
        assert all(cell.data_type != "f" for row in book.active for cell in row)
        book.close()
    else:
        rows = list(csv.reader(io.StringIO(exported.decode("utf-8-sig"))))
        assert rows[0] == STRUCTURED + AUDIT
        assert rows[1][:7] == ["'=SUM(1)", "'-Office", "active", "'+12345", "US", "''=literal", "'@City"]
    assert {key: parse(exported, "locations." + format)[0][key] for key in body} == body
    row = saved.json()
    assert api.post(f"{BASE}/{row['id']}/delete", headers=headers, json={"expected_version": 1}).status_code == 200
    filename = "locations." + format
    report = review(api, headers, exported, filename).json()
    assert report["valid"]
    # A changed city, filename or actor is not the reviewed operation, even if
    # the upload remains a syntactically valid current-format backup.
    changed = transfer_file([STRUCTURED, ["'=SUM(1)", "'-Office", "active", "'+12345", "US", "''=literal", "Changed"]], format)
    assert commit(api, headers, changed, report["digest"], filename).status_code == 409
    assert commit(api, headers, exported, report["digest"], "renamed." + format).status_code == 409
    from app.services.location_transfer import review_digest
    from app.services.auth import Identity
    session_id = db.scalar(select(AuditEvent.session_id).where(AuditEvent.action == "location_create"))
    current = Identity(actor, session_id=session_id)
    assert review_digest(current, exported, filename) == report["digest"]
    from types import SimpleNamespace
    other = SimpleNamespace(user=SimpleNamespace(id=uuid.uuid4()), session_id=session_id)
    assert review_digest(other, exported, filename) != report["digest"]
    assert review_digest(SimpleNamespace(user=actor, session_id="different-session"), exported, filename) != report["digest"]
    assert commit(api, headers, exported, report["digest"], filename).json() == {"imported": 1}
    restored = api.get(BASE, headers=headers).json()["items"][0]
    assert {key: restored[key] for key in body} == body


def test_eleven_column_workbook_bounds_include_hidden_coordinates():
    columns = STRUCTURED + AUDIT
    row = ["Room", "Office", "active", "411001", "IN", "Maharashtra", "Pune", "Ignored", "1900", "Ignored", "1900"]
    valid = workbook([columns] + [row] * 1000)
    assert len(parse(valid, "locations.xlsx")) == 1000
    for data in [
        workbook([columns + ["Extra"], row + ["hidden"]]),
        workbook([columns] + [row] * 1001),
        replace_zip(workbook([columns, row]), "xl/worksheets/sheet1.xml",
                    lambda xml: xml.replace(b'r="K2"', b'r="L2"')),
        replace_zip(workbook([columns, row]), "xl/worksheets/sheet1.xml",
                    lambda xml: xml.replace(b'r="K2"', b'r="K1002"')),
        replace_zip(workbook([columns, row]), "xl/worksheets/sheet1.xml",
                    lambda xml: xml.replace(b'ref="A1:K2"', b'ref="A1:C2"').replace(b'r="K2"', b'r="L2"')),
    ]:
        with pytest.raises(LocationError) as exc:
            parse(data, "locations.xlsx")
        assert exc.value.code == "location_invalid_file"


@pytest.mark.parametrize("grant", ["add", "edit", "delete", "import", "export"])
def test_location_pin_permissions_match_consuming_action_not_other_master(client, monkeypatch, grant):
    from test_zone_permissions import setup, staff_login, bearer
    from app.services import mr_postal
    api, db, _ = client
    admin, _, _, record, password = setup(api, db, [f"location.{grant}"])
    lookup = mr_postal.lookup
    calls = []
    monkeypatch.setattr(mr_postal, "_cache", OrderedDict())
    def provider(pin):
        calls.append(pin)
        return [dict(city="Pune", state="Maharashtra", country="India")]
    monkeypatch.setattr(mr_postal, "lookup", lambda db, actor, pin, **policy:
                        lookup(db, actor, pin, provider, **policy))
    assert api.get(BASE + "/postal/411001").status_code == 401
    staff = bearer(staff_login(api, record, password))
    for action in ("add", "edit"):
        response = api.get(BASE + "/postal/411001", params={"action": action}, headers=staff)
        assert response.status_code == (200 if grant == action else 403), response.text
        if grant == action:
            assert response.json()["choices"] == [dict(city="Pune", state="Maharashtra", country="India")]
    assert len(calls) == (1 if grant in ("add", "edit") else 0)
    for resource in ("mrs", "doctors", "patients"):
        assert api.get(f"/api/v1/admin/{resource}/postal/411001", headers=staff).status_code == 403
    assert api.get(BASE + "/postal/411001?action=delete", headers=admin).status_code == 422


def test_location_pin_provider_failure_leaves_saved_geography_unchanged(client, monkeypatch):
    from app.services import mr_postal
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    body = dict(name="Retained geography", address="Office", status="active", **PUNE)
    saved = api.post(BASE, headers=headers, json=body).json()
    lookup = mr_postal.lookup
    monkeypatch.setattr(mr_postal, "_cache", OrderedDict())
    def unavailable(pin):
        raise TimeoutError("Synthetic provider failure")
    monkeypatch.setattr(mr_postal, "lookup", lambda db, actor, pin, **policy:
                        lookup(db, actor, pin, unavailable, **policy))
    response = api.get(BASE + "/postal/411001?action=edit", headers=headers)
    assert response.status_code == 503
    assert response.json()["error"]["code"] == "mr_pin_unavailable"
    assert api.get(f"{BASE}/{saved['id']}", headers=headers).json() == saved


def test_unknown_indian_state_is_rejected_but_foreign_state_is_free_text(client):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    for country, valid in (("IN", False), ("US", True)):
        body = dict(name="Unknown state " + country, address="Office", status="active",
                    pincode="", countryCode=country, state="Not a recognized Indian state", city="City")
        assert api.post(BASE, headers=headers, json=body).status_code == (201 if valid else 422)
        data = transfer_file([STRUCTURED, [body[key] for key in ("name", "address", "status", *EMPTY_GEO)]], "csv")
        parsed = parse(data, "locations.csv")[0]
        assert bool(parsed["errors"]) is not valid


def add(api, headers, name="Supply room", status="active", address="Building A"):
    result = api.post(BASE, headers=headers, json={"name": name, "address": address, "status": status})
    assert result.status_code == 201, result.text
    return result.json()


def review(api, headers, data, filename="locations.csv"):
    return api.post(BASE + "/import/review", headers={**headers, "Content-Type": "application/octet-stream"},
                    params={"filename": filename}, content=data)


def commit(api, headers, data, digest, filename="locations.csv"):
    return api.post(BASE + "/import/commit", headers={**headers, "Content-Type": "application/octet-stream"},
                    params={"filename": filename, "digest": digest, "confirm": "true"}, content=data)


def test_persistence_filters_versions_creator_and_soft_delete(client):
    api, db, _ = client
    headers, actor = admin_headers(api, db)
    row = add(api, headers, "  City \t Dispatch  ")
    assert row["name"] == "City Dispatch"
    assert row["createdBy"] == row["updatedBy"] == "Super Admin"
    assert row["createdAt"] == row["updatedAt"]
    assert api.post(BASE, headers=headers, json={"name": "CITY   dispatch", "address": "Building A", "status": "inactive"}).status_code == 409
    edited = api.post(f"{BASE}/{row['id']}/edit", headers=headers,
                     json={"name": "City revised", "address": "Building B", "status": "active", "expected_version": 1}).json()
    assert edited["createdAt"] == row["createdAt"] and edited["version"] == 2
    assert edited["updatedAt"] > row["updatedAt"]
    assert edited["address"] == "Building B"
    assert api.get(BASE, headers=headers, params={"query": "building b", "limit": 1}).json()["filtered"] == 1
    assert api.post(f"{BASE}/{row['id']}/status", headers=headers,
                    json={"status": "inactive", "expected_version": 1}).status_code == 409
    assert api.post(f"{BASE}/{row['id']}/status", headers=headers,
                    json={"status": "inactive", "expected_version": 2}).json()["version"] == 3
    assert api.post(BASE, headers=headers, json={"name": "CITY REVISED", "address": "Different", "status": "active"}).status_code == 409
    for name in ("East", "West", "South"):
        add(api, headers, name)
    result = api.get(BASE, headers=headers, params={"status": "inactive", "query": "city", "limit": 1}).json()
    assert result["total"] == 4 and result["filtered"] == 1 and result["items"][0]["id"] == row["id"]
    assert api.get(BASE, headers=headers, params={"query": "%"}).json()["filtered"] == 0
    first = api.get(BASE, headers=headers, params={"limit": 2}).json()["items"]
    second = api.get(BASE, headers=headers, params={"limit": 2, "offset": 2}).json()["items"]
    assert len({entry["id"] for entry in first + second}) == 4
    assert api.post(f"{BASE}/{row['id']}/delete", headers=headers, json={"expected_version": 2}).status_code == 409
    assert api.post(f"{BASE}/{row['id']}/delete", headers=headers, json={"expected_version": 3}).json()["version"] == 4
    db.expire_all()
    stored = db.get(StorageLocation, uuid.UUID(row["id"]))
    assert stored.deleted_at is not None and stored.deleted_by == actor.id
    assert stored.updated_at == stored.deleted_at and stored.updated_by == stored.created_by == actor.id
    deletion_time = stored.deleted_at
    assert api.post(f"{BASE}/{row['id']}/delete", headers=headers, json={"expected_version": 4}).status_code == 404
    assert api.post(f"{BASE}/{row['id']}/edit", headers=headers,
                    json={"name": "Changed tombstone", "address": "No", "status": "active", "expected_version": 4}).status_code == 404
    db.expire_all()
    assert db.get(StorageLocation, stored.id).deleted_at == deletion_time
    assert api.get(f"{BASE}/{row['id']}", headers=headers).status_code == 404
    assert api.get(BASE, headers=headers).json()["total"] == 3
    assert "City revised" not in api.get(BASE + "/export", headers=headers).text
    assert add(api, headers, "City revised")["id"] != row["id"]
    event = db.scalar(select(AuditEvent).where(AuditEvent.action == "location_delete"))
    assert event.actor_id == actor.id and event.resource_id == stored.id and event.session_id
    # Location operations do not alter the separate Zone table.
    assert api.get("/api/v1/admin/zones", headers=headers).json()["total"] == 0


def test_authorization_forged_audit_inputs_and_expired_auth(client):
    api, db, _ = client
    assert api.get(BASE).status_code == 401
    mr = create_user(db, "location-mr@example.com")
    headers = {"Authorization": "Bearer " + login(api, mr.email).json()["access_token"]}
    assert api.get(BASE, headers=headers).status_code == 403
    assert api.get(BASE + "/export", headers=headers).status_code == 403
    assert review(api, headers, b"Storage Location,Address,Status\nPrivate,Active").status_code == 403
    headers, _ = admin_headers(api, db)
    for extra in ({"createdBy": "Forged"}, {"created_by": str(uuid.uuid4())}, {"deleted_at": "now"},
                  {"deleted_by": str(uuid.uuid4())}, {"created_at": "1900"}, {"updated_at": "1900"},
                  {"updated_by": str(uuid.uuid4())}, {"version": 50}):
        assert api.post(BASE, headers=headers, json={"name": "Forged", "address": "Building A", "status": "active", **extra}).status_code == 422
    for name in ("", "  ", "x" * 201):
        assert api.post(BASE, headers=headers, json={"name": name, "address": "Building A", "status": "active"}).status_code == 422
    for address in ("", "  ", "x" * 2001, "Nul\x00address"):
        assert api.post(BASE, headers=headers, json={"name": "Bounded", "address": address, "status": "active"}).status_code == 422
    assert api.post(BASE, headers=headers, json={"name": "Missing address", "status": "active"}).status_code == 422
    api.post("/api/v1/auth/logout", headers={"Origin": "http://testserver"}, json={})
    assert api.post(BASE, headers=headers, json={"name": "Expired", "address": "Building A", "status": "active"}).status_code == 401


def test_review_confirm_revalidation_and_import_attribution(client):
    api, db, _ = client
    headers, actor = admin_headers(api, db)
    data = b"\xef\xbb\xbfStorage Location,Address,Status,Created By,Created At,Updated By,Updated At\nFirst,Building A,Active,Forged,1900,Forged,1900\nSecond,Building A,Inactive,Forged,1900,Forged,1900"
    report = review(api, headers, data).json()
    assert report["valid"] and db.scalar(select(func.count()).select_from(StorageLocation)) == 0
    assert commit(api, headers, data + b"\n", report["digest"]).status_code == 409
    assert commit(api, headers, data, "0" * 64).status_code == 409
    assert api.post(BASE + "/import/commit", headers=headers,
                    params={"filename": "locations.csv", "digest": report["digest"], "confirm": "false"}).status_code == 422
    add(api, headers, "Second")
    assert commit(api, headers, data, report["digest"]).status_code == 409
    assert not db.scalar(select(StorageLocation).where(StorageLocation.name == "First"))
    duplicates = review(api, headers, b"Storage Location,Address,Status\nCity   Dispatch,Building A,active\ncity dispatch,Building A,inactive").json()
    assert not duplicates["valid"] and duplicates["rows"][1]["errors"]
    invalid = review(api, headers, b"Storage Location,Address,Status\nBad\x00Name,Building A,active").json()
    assert not invalid["valid"] and invalid["rows"][0]["errors"]
    clean = b"Storage Location,Address,Status\nImported,Building A,active"
    report = review(api, headers, clean).json()
    # A new session cannot reuse the old session's review confirmation.
    renewed_headers = {"Authorization": "Bearer " + login(api, actor.email).json()["access_token"]}
    assert commit(api, renewed_headers, clean, report["digest"]).status_code == 409
    report = review(api, renewed_headers, clean).json()
    assert commit(api, renewed_headers, clean, report["digest"]).json() == {"imported": 1}
    stored = db.scalar(select(StorageLocation).where(StorageLocation.name == "Imported"))
    assert stored.created_by == stored.updated_by == actor.id
    audited = b"Storage Location,Address,Status,Created By,Created At,Updated By,Updated At\nAudit ignored,Office,active,Forged,1900,Forged,1900"
    report = review(api, renewed_headers, audited).json()
    assert commit(api, renewed_headers, audited, report["digest"]).status_code == 200
    stored = db.scalar(select(StorageLocation).where(StorageLocation.name == "Audit ignored"))
    assert stored.created_by == stored.updated_by == actor.id and stored.created_at.year > 1900
    assert stored.created_at == stored.updated_at


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_filtered_export_safe_roundtrip_and_audit_ignored(client, format):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    for name in ("=SUM(1)", "'=literal", "+dispatch", "Plain"):
        add(api, headers, name, address="'=street" if name == "Plain" else "@Office")
    add(api, headers, "Hidden inactive", "inactive")
    assert len(api.get(BASE, headers=headers, params={"limit": 2}).json()["items"]) == 2
    response = api.get(BASE + "/export", headers=headers, params={"status": "active", "format": format})
    assert response.status_code == 200
    assert response.headers["content-disposition"].endswith(f'.{format}"')
    assert response.headers["content-type"].startswith("text/csv" if format == "csv" else "application/vnd.openxml")
    if format == "xlsx":
        book = load_workbook(io.BytesIO(response.content))
        assert all(cell.data_type != "f" for row in book.active for cell in row)
        book.close()
    assert {row["name"] for row in parse(response.content, "locations." + format)} == {"=SUM(1)", "'=literal", "+dispatch", "Plain"}
    assert {row["address"] for row in parse(response.content, "locations." + format)} == {"'=street", "@Office"}
    for row in list(db.scalars(select(StorageLocation))):
        api.post(f"{BASE}/{row.id}/delete", headers=headers, json={"expected_version": row.version})
    report = review(api, headers, response.content, "locations." + format).json()
    assert report["valid"]
    assert commit(api, headers, response.content, report["digest"], "locations." + format).json()["imported"] == 4


def test_parser_security_and_upload_layers(client):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    valid = workbook([["Storage Location", "Address", "Status"], ["Workbook", "Building A", "Active"]])
    assert review(api, headers, valid, "locations.xlsx").json()["valid"]
    hostile = [
        (b"\xff", "locations.csv"), (b"Not a zip", "locations.xlsx"), (valid, "locations.xls"),
        (b"Wrong,Status\nName,active", "locations.csv"),
        (b'Storage Location,Address,Status\n"unterminated,active', "locations.csv"),
        (workbook([["Storage Location", "Address", "Status"], ["=1+1", "Building A", "Active"]]), "locations.xlsx"),
        (replace_zip(valid, "xl/worksheets/sheet1.xml", lambda x: x.replace(b'r="A2"', b'r="ZZ99999"')), "locations.xlsx"),
        (replace_zip(valid, "xl/_rels/workbook.xml.rels", lambda x: x.replace(b'Target=', b'TargetMode="External" Target=', 1)), "locations.xlsx"),
        (replace_zip(valid, "[Content_Types].xml", lambda x: x.replace(b'spreadsheetml.sheet.main+xml', b'spreadsheetml.macroEnabled.main+xml')), "locations.xlsx"),
        (replace_zip(valid, "xl/worksheets/sheet1.xml", lambda x: b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]>' + x), "locations.xlsx"),
        (workbook([["Storage Location", "Address", "Status"], ["x" * 10001, "Building A", "active"]]), "locations.xlsx"),
    ]
    for data, filename in hostile:
        assert review(api, headers, data, filename).status_code == 422
    assert review(api, headers, b"x" * (MAX_BYTES + 1)).status_code == 413
    assert review(api, headers, b"Storage Location,Address,Status\n" + b"Partner,Building A,Active\n" * 1001).status_code == 422
    assert api.post(BASE + "/import/review", headers={**headers, "Content-Type": "application/octet-stream"},
                    params={"filename": "locations.csv"}, content=iter([b"x" * 1024 * 1024] * 3)).status_code == 413
    assert api.post(BASE, headers=headers, content=b"x" * (1024 * 1024 + 1)).status_code == 413
    data = b"Storage Location,Address,Status\nMultipart,Building A,active"
    assert api.post(BASE + "/import/review", headers=headers, params={"filename": "locations.csv"},
                    files={"file": ("locations.csv", data, "text/csv")}).json()["valid"]
    assert api.post(BASE + "/import/review", headers=headers, params={"filename": "locations.csv"},
                    files={"file": ("locations.csv", b"x" * (MAX_BYTES + 1))}).status_code == 413
    assert api.post(BASE + "/import/review", headers=headers, params={"filename": "locations.csv"},
                    files=[("file", ("a.csv", data)), ("file", ("b.csv", data))]).status_code == 422
    for header in ("Storage Location,Address,Status,ID", "Storage Location,Address,Status,version",
                   "Storage Location,Address,Status,deleted_at", "Storage Location ,Address,Status"):
        assert review(api, headers, (header + "\nName,Office,active,ignored").encode()).status_code == 422
    assert not review(api, headers, b"Storage Location,Address,Status\nName,,active").json()["valid"]
    # The global request layer and endpoint extraction layer preserve both bounds.
    from app.core.request_limits import body_limit
    assert body_limit("POST", BASE + "/import/review", 10) == MAX_BYTES + 65536
    assert body_limit("POST", "/api/v1/files", 10) == 10
    assert body_limit("POST", BASE, 10) == 1048576
    assert api.post(BASE + "/import/review", headers={**headers, "Content-Type": "application/octet-stream",
                    "Content-Length": str(MAX_BYTES + 65537)}, params={"filename": "locations.csv"}, content=b"x").status_code == 413
    assert api.post(BASE + "/import/review", headers=headers, params={"filename": "locations.csv"},
                    files={"file": ("x" * 66000 + ".csv", data)}).status_code == 422


def test_export_5000_bound_is_explicit(client):
    api, db, _ = client
    headers, actor = admin_headers(api, db)
    db.add_all([StorageLocation(name=f"Bound {index}", address="Building A", status="active", created_by=actor.id, updated_by=actor.id)
                for index in range(5001)])
    db.commit()
    assert api.get(BASE + "/export", headers=headers).json()["error"]["code"] == "location_export_limit"
    assert api.get(BASE + "/export", headers=headers, params={"query": "Bound 5000"}).status_code == 200
    assert api.get(BASE + "/export", headers=headers, params={"query": "Bound 1"}).text.count("\r\n") > 1000
