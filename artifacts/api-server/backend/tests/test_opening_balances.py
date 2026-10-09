"""Disposable authenticated API coverage. No configured accounts/databases."""
import csv
import io
import uuid
import zipfile
from decimal import Decimal
import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from app.db.models import AuditEvent
from app.db.opening_balance_models import OpeningBalance, OpeningBalanceImportReview
from app.services import opening_balance_transfer as transfer
from test_sessions import client, create_user, login
from test_doctors import setup_doctor, fields as doctor_fields, add as add_doctor

BASE = "/api/v1/admin/opening-balances"


def setup_balance(api, db):
    headers, actor, mr, zone = setup_doctor(api, db)
    doctor = add_doctor(api, headers, doctor_fields(mr))
    return headers, actor, doctor


def body(doctor, **changes):
    return {**dict(startYear=2026, endYear=2027, doctorId=doctor["id"], amount="-9999999999999.99", status="active"), **changes}


def add(api, headers, doctor, **changes):
    response = api.post(BASE, headers=headers, json=body(doctor, **changes))
    assert response.status_code == 201, response.text
    return response.json()


def upload(api, headers, data, filename="balance.csv", digest=None):
    return api.post(BASE + ("/import/commit" if digest else "/import/review"),
                    headers={**headers, "Content-Type": "application/octet-stream"}, content=data,
                    params={"filename": filename, **({"confirm": True, "digest": digest} if digest else {})})


def file(doctor, *rows, full=False):
    stream = io.StringIO()
    writer = csv.writer(stream)
    writer.writerow(transfer.HEADERS if full else transfer.HEADERS[:5])
    for start, amount, status in rows or [(2026, "-25.50", "active")]:
        writer.writerow([start, start + 1, doctor["registrationNumber"], amount, status] +
                        (["UNTRUSTED", "1900-01-01", "UNTRUSTED", "1900-01-02"] if full else []))
    return stream.getvalue().encode()


def test_crud_exact_audit_filter_versions_and_soft_delete(client):
    api, db, _ = client
    headers, actor, doctor = setup_balance(api, db)
    row = add(api, headers, doctor)
    assert row["amount"] == "-9999999999999.99" and row["createdBy"] == "Super Admin"
    assert row["doctorName"] == doctor["name"] and row["registrationNumber"] == doctor["registrationNumber"]
    assert api.get(BASE + "/" + row["id"], headers=headers).json() == row
    assert api.post(BASE, headers=headers, json=body(doctor)).status_code == 409
    edit = api.post(BASE + "/" + row["id"] + "/edit", headers=headers,
                   json={**body(doctor, amount="+0.00"), "expected_version": 1})
    assert edit.status_code == 200, edit.text
    assert edit.json()["amount"] == "0.00" and edit.json()["version"] == 2
    assert edit.json()["createdAt"] == row["createdAt"]
    assert api.post(BASE + "/" + row["id"] + "/status", headers=headers,
                    json={"status": "inactive", "expected_version": 1}).status_code == 409
    changed = api.post(BASE + "/" + row["id"] + "/status", headers=headers,
                       json={"status": "inactive", "expected_version": 2})
    assert changed.status_code == 200
    for query in ("2026-2027", "2026–2027", doctor["name"], doctor["registrationNumber"], "0.00"):
        page = api.get(BASE, headers=headers, params={"query": query, "status": "inactive", "limit": 1}).json()
        assert page["filtered"] is None and page["partial"] and page["total"] == 1 and len(page["items"]) == 1
    literal = api.get(BASE, headers=headers, params={"query": "%"}).json()
    assert literal["filtered"] is None and literal["partial"] and not literal["items"]
    assert api.get(BASE, headers=headers, params={"status": "active"}).json()["filtered"] == 0
    deleted = api.post(BASE + "/" + row["id"] + "/delete", headers=headers, json={"expected_version": 3})
    assert deleted.status_code == 200 and deleted.json()["version"] == 4
    assert api.get(BASE, headers=headers).json()["total"] == 0
    assert api.get(BASE + "/" + row["id"], headers=headers).status_code == 404
    db.expire_all()
    historical = db.get(OpeningBalance, uuid.UUID(row["id"]))
    assert historical.deleted_by == actor.id and historical.created_by == actor.id
    assert historical.deleted_at is not None and historical.amount == Decimal("0.00")
    assert db.scalar(select(func.count()).select_from(AuditEvent).where(
        AuditEvent.resource_id == historical.id, AuditEvent.actor_id == actor.id)) == 4
    assert add(api, headers, doctor, amount="9999999999999.99")["id"] != row["id"]
    searched = api.get(BASE, headers=headers, params={"query": "99,99,99,99,99,999.99"}).json()
    assert searched["filtered"] is None and len(searched["items"]) == 1


@pytest.mark.parametrize("changes", [
    {"amount": 0.1}, {"amount": "1.234"}, {"amount": "1.000"}, {"amount": "1e2"},
    {"amount": "9999999999999.991"}, {"amount": "10000000000000"}, {"amount": "NaN"},
    {"amount": "-10000000000000"}, {"startYear": 1899, "endYear": 1900},
    {"startYear": 9999, "endYear": 10000}, {"startYear": "2026"}, {"endYear": 2028},
    {"createdBy": "spoof"}, {"status": "unknown"}, {"doctorId": "sample-doctor-1"},
])
def test_validation_no_rounding_or_attribution_injection(client, changes):
    api, db, _ = client
    headers, _, doctor = setup_balance(api, db)
    response = api.post(BASE, headers=headers, json=body(doctor, **changes))
    assert response.status_code == 422, response.text
    assert api.get(BASE, headers=headers).json()["total"] == 0


def test_minimal_bounded_choices_saved_inactive_retention(client):
    api, db, _ = client
    headers, _, doctor = setup_balance(api, db)
    row = add(api, headers, doctor)
    choices = api.get(BASE + "/references", headers=headers, params={"limit": 1}).json()
    assert choices["total"] == 1 and set(choices["items"][0]) == {"id", "name", "registrationNumber", "usable"}
    assert "phone" not in choices["items"][0]
    response = api.post("/api/v1/admin/doctors/" + doctor["id"] + "/status", headers=headers,
                        json={"status": "inactive", "expected_version": 1})
    assert response.status_code == 200
    assert api.get(BASE + "/references", headers=headers).json()["items"] == []
    saved = api.get(BASE + "/references", headers=headers, params={"balance_id": row["id"]}).json()
    assert saved["total"] == 0 and saved["items"][0]["usable"] is False
    assert api.post(BASE, headers=headers, json=body(doctor, startYear=2027, endYear=2028)).status_code == 409
    edit = api.post(BASE + "/" + row["id"] + "/edit", headers=headers,
                   json={**body(doctor, amount="-25.50"), "expected_version": 1})
    assert edit.status_code == 200 and edit.json()["amount"] == "-25.50"


def test_import_review_binding_reference_changes_consumption_and_atomic_batch(client):
    api, db, _ = client
    headers, _, doctor = setup_balance(api, db)
    data = file(doctor, (1900, "-25.50", "active"), (9998, "0", "inactive"), full=True)
    review = upload(api, headers, data)
    assert review.status_code == 200 and review.json()["valid"]
    assert api.get(BASE, headers=headers).json()["total"] == 0
    assert upload(api, headers, data + b"\n", digest=review.json()["digest"]).status_code == 409
    # A failed commit consumes its previous review.
    assert upload(api, headers, data, digest=review.json()["digest"]).status_code == 409
    review = upload(api, headers, data).json()
    api.post("/api/v1/admin/doctors/" + doctor["id"] + "/status", headers=headers,
             json={"status": "inactive", "expected_version": 1})
    assert upload(api, headers, data, digest=review["digest"]).status_code == 409
    assert api.get(BASE, headers=headers).json()["total"] == 0
    api.post("/api/v1/admin/doctors/" + doctor["id"] + "/status", headers=headers,
             json={"status": "active", "expected_version": 2})
    assert upload(api, headers, data, digest=review["digest"]).status_code == 409
    review = upload(api, headers, data).json()
    imported = upload(api, headers, data, digest=review["digest"])
    assert imported.status_code == 200 and imported.json() == {"imported": 2}, imported.text
    assert upload(api, headers, data, digest=review["digest"]).status_code == 409
    page = api.get(BASE, headers=headers).json()
    assert page["total"] == 2 and all(r["createdBy"] == "Super Admin" for r in page["items"])
    duplicate = upload(api, headers, file(doctor, (2027, "1", "active"), (2027, "2", "inactive"))).json()
    assert not duplicate["valid"] and duplicate["rows"][1]["errors"]
    missing = upload(api, headers, file({**doctor, "registrationNumber": "SAMPLE-LOCAL"})).json()
    assert not missing["valid"] and missing["rows"][0]["errors"]


def test_csv_xlsx_filtered_exports_roundtrip_text_money_and_durable_downloads(client):
    api, db, _ = client
    headers, _, doctor = setup_balance(api, db)
    row = add(api, headers, doctor)
    add(api, headers, doctor, startYear=2027, endYear=2028, amount="5.20", status="inactive")
    for fmt in ("csv", "xlsx"):
        # Omitted caller key still requires durable evidence (server assigns it).
        assert api.get(BASE + "/sample", headers=headers, params={"format": fmt}).headers["X-Download-Log"]
        sample = api.get(BASE + "/sample", headers={**headers, "X-Download-Initiation": str(uuid.uuid4())}, params={"format": fmt})
        assert sample.status_code == 200 and sample.headers["X-Download-Log"]
        export = api.get(BASE + "/export", headers={**headers, "X-Download-Initiation": str(uuid.uuid4())},
                         params={"format": fmt, "status": "active", "query": doctor["registrationNumber"]})
        assert export.status_code == 200 and export.headers["X-Download-Log"]
        assert "session" not in export.content.decode(errors="ignore").lower()
        if fmt == "xlsx":
            book = load_workbook(io.BytesIO(export.content))
            assert book.active["D2"].data_type == "s"
            assert book.active.max_row == 2
            book.close()
        api.post(BASE + "/" + row["id"] + "/delete", headers=headers, json={"expected_version": row["version"]})
        reviewed = upload(api, headers, export.content, "export." + fmt).json()
        assert reviewed["valid"], reviewed
        assert reviewed["rows"][0]["amount"] == "-9999999999999.99"
        assert upload(api, headers, export.content, "export." + fmt, reviewed["digest"]).json()["imported"] == 1
        row = api.get(BASE, headers=headers, params={"status": "active"}).json()["items"][0]


def test_original_workbook_decimal_lexemes_and_inert_limits(client):
    api, db, _ = client
    headers, _, doctor = setup_balance(api, db)
    output = io.BytesIO()
    book = Workbook()
    book.active.append(transfer.HEADERS[:5])
    book.active.append([2026, 2027, doctor["registrationNumber"], -1.25, "active"])
    book.save(output)
    book.close()
    def changed(lexeme):
        result = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(output.getvalue())) as source, zipfile.ZipFile(result, "w") as target:
            for name in source.namelist():
                data = source.read(name)
                if name == "xl/worksheets/sheet1.xml":
                    data = data.replace(b"<v>-1.25</v>", f"<v>{lexeme}</v>".encode())
                target.writestr(name, data)
        return result.getvalue()
    for lexeme, valid in [("-9999999999999.99", True), ("-9.99999999999999E12", True),
                          ("1.001", False), ("9999999999999.991", False), ("1E99", False)]:
        report = upload(api, headers, changed(lexeme), "exact.xlsx")
        assert report.status_code == 200 and report.json()["valid"] == valid, report.text
        if valid:
            assert report.json()["rows"][0]["amount"] == "-9999999999999.99"
    assert upload(api, headers, b"x" * (2 * 1024 * 1024 + 1)).status_code == 413
    assert upload(api, headers, file(doctor, *[(2026, "0", "active")] * 1001)).status_code == 422
    assert upload(api, headers, b"fake workbook", "fake.xlsx").status_code == 422
    assert upload(api, headers, b"headers\n=HYPERLINK(\"https://bad\")", "bad.csv").status_code == 422


def test_nonprotected_identity_has_no_workflow_or_reference_access(client):
    api, db, _ = client
    headers, _, doctor = setup_balance(api, db)
    user = create_user(db, "balance-mr@example.com")
    response = login(api, user.email)
    other = {"Authorization": "Bearer " + response.json()["access_token"]}
    for path in ("", "/references", "/sample", "/export"):
        assert api.get(BASE + path, headers=other).status_code == 403
    assert api.post(BASE, headers=other, json=body(doctor)).status_code == 403
    assert upload(api, other, file(doctor)).status_code == 403


def test_bounded_doctor_paging_and_export_cap(client, monkeypatch):
    from app.db.doctor_models import DoctorDirectory
    from app.schemas.doctors import DoctorFields
    from app.core.security import utcnow
    api, db, _ = client
    headers, actor, mr, _zone = setup_doctor(api, db)
    now = utcnow()
    docs = [DoctorDirectory(**DoctorFields(**doctor_fields(mr, name=f"Paged {i:03d}",
            registrationNumber=f"PAGED-{i:03d}")).model_dump(), version=1, verification="unverified",
            created_by=actor.id, updated_by=actor.id, created_at=now, updated_at=now) for i in range(55)]
    db.add_all(docs); db.commit()
    page1 = api.get(BASE + "/references", headers=headers, params={"query": "Paged", "limit": 50}).json()
    page2 = api.get(BASE + "/references", headers=headers, params={"query": "Paged", "limit": 50, "cursor": page1["nextCursor"]}).json()
    assert len(page1["items"]) == 50 and len(page2["items"]) == 5
    assert page1["total"] is None and page2["total"] is None
    assert page1["partial"] and page2["partial"]
    assert not ({r["id"] for r in page1["items"]} & {r["id"] for r in page2["items"]})
    assert api.get(BASE + "/references", headers=headers, params={"limit": 101}).status_code == 422
    db.add_all([OpeningBalance(startYear=1900+i, endYear=1901+i, doctorId=docs[0].id, amount=Decimal("0.00"),
                              status="active", version=1, created_by=actor.id, updated_by=actor.id,
                              created_at=now, updated_at=now) for i in range(5001)])
    db.commit()
    oversized = api.get(BASE + "/export", headers=headers, params={"format": "csv"})
    assert oversized.status_code == 409 and "nothing downloaded" in oversized.text
    page = api.get(BASE, headers=headers, params={"limit": 10, "offset": 5000}).json()
    assert page["total"] == page["filtered"] == 5001 and len(page["items"]) == 1


def test_logging_failure_and_invalid_replacement_do_not_release_or_replay(client, monkeypatch):
    from fastapi import HTTPException
    from app.api.v1 import opening_balances as routes
    api, db, _ = client
    headers, _, doctor = setup_balance(api, db)
    data = file(doctor)
    reviewed = upload(api, headers, data).json()
    assert upload(api, headers, b"bad replacement", "bad.csv").status_code == 422
    assert upload(api, headers, data, digest=reviewed["digest"]).status_code == 409
    def fail(*_args, **_kwargs):
        raise HTTPException(503, "Download logging unavailable")
    monkeypatch.setattr(routes, "server_record", fail)
    for path in ("/sample", "/export"):
        result = api.get(BASE + path, headers=headers)
        assert result.status_code == 503 and not result.headers.get("Content-Disposition")
