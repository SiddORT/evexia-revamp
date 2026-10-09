"""Synthetic Allergen API/transfer fixtures; never managed data or credentials."""
import csv
import io
import uuid
from datetime import timedelta
from decimal import Decimal
import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from app.core.security import utcnow
from app.db.models import AuditEvent
from app.db.allergen_models import AllergenProduct
from app.db.download_models import DownloadLog
from app.db.product_category_models import ProductCategory
from app.db.location_models import StorageLocation
from app.services import allergens as service, allergen_transfer as transfer
from app.schemas.allergens import AllergenFields
from app.services.auth import identity_from_token
from test_sessions import client, create_user, login
from test_reporting import admin_headers
from test_zone_permissions import setup, staff_login, bearer
from test_zones import workbook

BASE = "/api/v1/admin/allergens"


def references(api, headers):
    c = api.post("/api/v1/admin/product-categories", headers=headers, json={
        "name": "Shared category %", "unit_price": "0", "description": "", "status": "active"})
    l = api.post("/api/v1/admin/storage-locations", headers=headers, json={
        "name": "Shared location _", "address": "Laboratory", "status": "active"})
    assert c.status_code == l.status_code == 201, (c.text, l.text)
    return c.json(), l.json()


def fields(c, l, **values):
    return {**dict(name="Diagnostic reagent", category_id=c["id"], storage_location_id=l["id"],
                   selling_price=None, gst="12", concentration="10 mg/mL", threshold_limit=None,
                   status="active", mix=False), **values}


def add(api, headers, c, l, **values):
    result = api.post(BASE, headers=headers, json=fields(c, l, **values))
    assert result.status_code == 201, result.text
    return result.json()


def review(api, headers, data, filename="allergens.csv"):
    return api.post(BASE + "/import/review", headers={**headers, "Content-Type": "application/octet-stream"},
                    params={"filename": filename}, content=data)


def commit(api, headers, data, digest, filename="allergens.csv"):
    return api.post(BASE + "/import/commit", headers={**headers, "Content-Type": "application/octet-stream"},
                    params={"filename": filename, "digest": digest, "confirm": "true"}, content=data)


def data(c, l, name="Import reagent", legacy=False, mix="Mix"):
    cells = [name, c["name"], "999999999999.999999", "100.000000", l["name"], "10 mg/mL", "0", "inactive", mix]
    if legacy:
        cells.insert(7, "HSN deliberately discarded")
        cells[-1] = "Allergens" if mix == "Mix" else "No Mix"
    return transfer.encode([transfer.LEGACY if legacy else transfer.HEADERS[:9], cells], "csv")


def test_crud_versions_audit_and_soft_deletion(client):
    api, db, _ = client
    headers, actor = admin_headers(api, db)
    c, l = references(api, headers)
    assert api.get(BASE, headers=headers).json()["total"] == 0
    row = add(api, headers, c, l, name="  Diagnostic   reagent ", selling_price="0", mix=True)
    assert row["mix"] is True and row["selling_price"] == "0.000000" and row["threshold_limit"] is None
    assert row["name"] == "Diagnostic reagent"
    assert row["createdBy"] == row["updatedBy"] == "Super Admin"
    assert row["createdAt"] == row["updatedAt"]
    assert api.get(BASE + "/" + row["id"], headers=headers).json() == row
    assert api.post(BASE, headers=headers, json=fields(c, l, name="DIAGNOSTIC reagent")).status_code == 409
    edited = api.post(f"{BASE}/{row['id']}/edit", headers=headers, json={
        **fields(c, l, name="Changed", mix=False, selling_price="999999999999.999999"), "expected_version": 1})
    assert edited.status_code == 200, edited.text
    changed = edited.json()
    assert changed["mix"] is False and changed["version"] == 2
    assert changed["createdAt"] == row["createdAt"] and changed["updatedAt"] > row["updatedAt"]
    for suffix, body in (("edit", {**fields(c, l), "expected_version": 1}),
                         ("status", {"status": "inactive", "expected_version": 1}), ("delete", {"expected_version": 1})):
        assert api.post(f"{BASE}/{row['id']}/{suffix}", headers=headers, json=body).status_code == 409
    assert api.post(f"{BASE}/{row['id']}/status", headers=headers, json={"status": "inactive", "expected_version": 2}).json()["version"] == 3
    assert api.post(f"{BASE}/{row['id']}/delete", headers=headers, json={"expected_version": 3}).json()["version"] == 4
    assert api.get(BASE, headers=headers).json()["total"] == 0
    assert api.get(f"{BASE}/{row['id']}", headers=headers).status_code == 404
    db.expire_all()
    deleted = db.get(AllergenProduct, uuid.UUID(row["id"]))
    assert deleted.deleted_by == deleted.updated_by == actor.id and deleted.deleted_at == deleted.updated_at
    add(api, headers, c, l, name="changed")
    evidence = list(db.scalars(select(AuditEvent).where(AuditEvent.resource_type == "allergen")))
    assert {"allergen_create", "allergen_edit", "allergen_status", "allergen_delete"} <= {e.action for e in evidence}
    assert all(e.actor_id == actor.id and e.session_id and e.request_id for e in evidence)


@pytest.mark.parametrize("key,value", [
    ("name", ""), ("name", "x" * 201), ("concentration", ""), ("concentration", "x" * 201),
    ("concentration", "a\0b"), ("selling_price", -1), ("selling_price", 0), ("selling_price", "1e2"),
    ("selling_price", "0.0000001"), ("selling_price", "1000000000000"), ("selling_price", ""),
    ("gst", "100.000001"), ("gst", "-1"), ("gst", None), ("gst", True), ("gst", "NaN"),
    ("threshold_limit", "-1"), ("threshold_limit", "1.1234567"), ("mix", "false"), ("mix", 0),
    ("status", "unknown"), ("hsnCode", "3822"), ("createdBy", "forged"), ("version", 99),
])
def test_invalid_fields_no_rounding_or_forged_attribution(client, key, value):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    c, l = references(api, headers)
    result = api.post(BASE, headers=headers, json=fields(c, l, **{key: value}))
    assert result.status_code == 422, result.text
    assert db.scalar(select(func.count()).select_from(AllergenProduct)) == 0
    assert "3822" not in result.text and "forged" not in result.text


def test_reference_choices_retention_and_save_time_changes(client):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    c, l = references(api, headers)
    row = add(api, headers, c, l)
    for kind in ("categories", "locations"):
        page = api.get(BASE + "/references/" + kind, headers=headers, params={"query": "%", "limit": 1}).json()
        assert page["total"] == (1 if kind == "categories" else 0)
    api.post(f"/api/v1/admin/product-categories/{c['id']}/status", headers=headers,
             json={"expected_version": 1, "status": "inactive"})
    api.post(f"/api/v1/admin/storage-locations/{l['id']}/delete", headers=headers, json={"expected_version": 1})
    assert api.get(BASE + "/references/categories", headers=headers).json()["total"] == 0
    assert api.get(BASE + "/references/locations", headers=headers).json()["total"] == 0
    detail = api.get(f"{BASE}/{row['id']}", headers=headers).json()
    assert detail["category_name"] == c["name"] and detail["category_status"] == "inactive"
    assert detail["storage_location_name"] == l["name"] and detail["storage_location_status"] == "deleted"
    assert api.post(BASE, headers=headers, json=fields(c, l, name="Not eligible")).status_code == 409
    assert api.post(f"{BASE}/{row['id']}/edit", headers=headers, json={
        **fields(c, l, name="Retained"), "expected_version": 1}).status_code == 200
    assert api.get(BASE + "/references/locations", headers=headers, params={"include_unusable": True}).json()["items"][0]["status"] == "deleted"
    assert api.post(f"{BASE}/{row['id']}/edit", headers=headers, json={
        **fields(c, l, category_id=str(uuid.uuid4())), "expected_version": 2}).status_code == 409


def test_all_combinable_filters_literal_search_pagination_null_vs_zero(client):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    c, l = references(api, headers)
    for index, price in enumerate((None, "0", "0.000001", "999999999999.999999")):
        add(api, headers, c, l, name=f"Product {index}", selling_price=price, mix=index % 2 == 1)
    add(api, headers, c, l, name="Inactive", status="inactive", mix=True, concentration="Unique concentrate")
    params = {"query": "%", "status": "active", "category_id": c["id"], "storage_location_id": l["id"],
              "mix": "mix", "min_price": "0", "max_price": "999999999999.999999", "limit": 1}
    result = api.get(BASE, headers=headers, params=params).json()
    assert result["total"] == 5 and result["filtered"] == 2 and len(result["items"]) == 1
    other = api.get(BASE, headers=headers, params={**params, "offset": 1}).json()
    assert other["items"][0]["id"] != result["items"][0]["id"]
    for fmt in ("csv", "xlsx"):
        exported = api.get(BASE + "/export", headers=headers, params={**params, "format": fmt})
        assert exported.status_code == 200 and len(transfer.parse(exported.content, f"backup.{fmt}")) == 2
    assert api.get(BASE, headers=headers, params={"min_price": "0", "max_price": "0"}).json()["filtered"] == 1
    assert api.get(BASE, headers=headers, params={"query": "_"}).json()["filtered"] == 5
    assert api.get(BASE, headers=headers, params={"query": "Unique concentrate"}).json()["filtered"] == 1
    assert api.get(BASE, headers=headers, params={"query": "' OR TRUE --"}).json()["filtered"] == 0
    assert api.get(BASE, headers=headers, params={"min_price": "", "max_price": ""}).json()["filtered"] == 5
    for bounds in ({"min_price": "1", "max_price": "0"}, {"max_price": "-1"}, {"min_price": "NaN"}, {"min_price": "1e2"}):
        for suffix in ("", "/export"):
            assert api.get(BASE + suffix, headers=headers, params=bounds).status_code == 422


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
@pytest.mark.parametrize("mix", [True, False])
def test_exact_full_field_export_reimport_and_durable_downloads(client, fmt, mix):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    c, l = references(api, headers)
    row = add(api, headers, c, l, name="=SUM(1,2)", selling_price="999999999999.999999",
              gst="99.123456", threshold_limit="0.000001", concentration="'quoted", mix=mix)
    initiation = str(uuid.uuid4())
    extra = {**headers, "X-Download-Initiation": initiation}
    exported = api.get(BASE + "/export", headers=extra, params={"format": fmt})
    assert exported.status_code == 200, exported.text
    assert exported.headers["X-Download-Log"] and exported.headers["Cache-Control"] == "no-store"
    assert headers["Authorization"] not in str(exported.content)
    assert b"HSN" not in exported.content
    assert api.get(BASE + "/export", headers=extra, params={"format": fmt}).headers["X-Download-Log"] == exported.headers["X-Download-Log"]
    assert api.get(BASE + "/sample", headers=extra, params={"format": fmt}).status_code == 409
    db.expire_all()
    assert db.scalar(select(func.count()).select_from(DownloadLog).where(DownloadLog.source == "allergen")) == 1
    parsed = transfer.parse(exported.content, f"backup.{fmt}")[0]
    assert not parsed["errors"] and parsed["selling_price"] == "999999999999.999999"
    assert parsed["concentration"] == "'quoted" and parsed["mix"] == ("Mix" if mix else "No Mix")
    if fmt == "xlsx":
        book = load_workbook(io.BytesIO(exported.content), data_only=False)
        assert book.active["C2"].data_type == "s" and book.active["A2"].data_type == "s"
        assert book.active["C2"].value == "999999999999.999999"
        book.close()
    report = review(api, headers, exported.content, f"backup.{fmt}").json()
    assert not report["valid"]
    api.post(f"{BASE}/{row['id']}/delete", headers=headers, json={"expected_version": 1})
    report = review(api, headers, exported.content, f"backup.{fmt}").json()
    assert report["valid"] and api.get(BASE, headers=headers).json()["total"] == 0
    assert commit(api, headers, exported.content, report["digest"], f"backup.{fmt}").json()["imported"] == 1
    imported = api.get(BASE, headers=headers).json()["items"][0]
    assert imported["id"] != row["id"] and imported["createdAt"] != row["createdAt"] and imported["version"] == 1
    assert imported["selling_price"] == row["selling_price"] and imported["gst"] == row["gst"] and imported["mix"] is mix
    for sample_fmt in ("csv", "xlsx"):
        sample = api.get(BASE + "/sample", headers=headers, params={"format": sample_fmt})
        assert sample.status_code == 200 and transfer.parse(sample.content, f"sample.{sample_fmt}")


@pytest.mark.parametrize("mix", ["Mix", "No Mix"])
def test_explicit_legacy_csv_mapping_ignores_hsn_and_has_no_local_id_fallback(client, mix):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    c, l = references(api, headers)
    file = data(c, l, legacy=True, mix=mix)
    report = review(api, headers, file).json()
    assert report["valid"] and report["rows"][0]["mix"] == mix
    assert "hsn" not in str(report).lower()
    assert commit(api, headers, file, report["digest"]).json()["imported"] == 1
    assert api.get(BASE, headers=headers).json()["items"][0]["mix"] is (mix == "Mix")
    unknown = data({**c, "name": "sample-category-1"}, l, name="Unknown", legacy=True)
    assert not review(api, headers, unknown).json()["valid"]


def test_review_file_session_reference_expiry_uniqueness_and_atomicity(client, monkeypatch):
    api, db, settings = client
    headers, _ = admin_headers(api, db)
    c, l = references(api, headers)
    file = data(c, l)
    report = review(api, headers, file).json()
    assert report["valid"] and db.scalar(select(func.count()).select_from(AllergenProduct)) == 0
    assert commit(api, headers, file + b"\n", report["digest"]).status_code == 409
    assert commit(api, headers, file, report["digest"], "renamed.csv").status_code == 409
    # A deterministic lease boundary, not a flaky real-time wait.
    future = utcnow() + timedelta(seconds=transfer.REVIEW_SECONDS + 1)
    with monkeypatch.context() as patch:
        patch.setattr(transfer, "utcnow", lambda: future)
        assert commit(api, headers, file, report["digest"]).status_code == 409
    api.post(f"/api/v1/admin/product-categories/{c['id']}/status", headers=headers, json={"expected_version": 1, "status": "inactive"})
    assert commit(api, headers, file, report["digest"]).status_code == 409
    assert not review(api, headers, file).json()["valid"]
    api.post(f"/api/v1/admin/product-categories/{c['id']}/status", headers=headers, json={"expected_version": 2, "status": "active"})
    # Reactivating a reference cannot resurrect an earlier review lease.
    assert commit(api, headers, file, report["digest"]).status_code == 409
    report = review(api, headers, file).json()
    add(api, headers, c, l, name="Import reagent")
    assert commit(api, headers, file, report["digest"]).status_code == 409
    two = transfer.encode([transfer.HEADERS[:9],
        ["Batch one", c["name"], "", "0", l["name"], "1:10", "", "active", "Mix"],
        ["Batch two", c["name"], "", "0", l["name"], "1:10", "", "active", "No Mix"]], "csv")
    report = review(api, headers, two).json()
    original = service.insert
    calls = 0
    def fail_after_first(db, actor, body):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise service.AllergenError("Synthetic failure", 409, "allergen_import_conflict")
        return original(db, actor, body)
    monkeypatch.setattr(service, "insert", fail_after_first)
    assert commit(api, headers, two, report["digest"]).status_code == 409
    assert api.get(BASE, headers=headers).json()["total"] == 1


def test_denied_identities_and_direct_service_revalidation(client):
    api, db, settings = client
    assert api.get(BASE).status_code == 401
    mr = create_user(db, "syntheticallergenmr@example.com")
    mr_token = login(api, mr.email).json()["access_token"]
    _, _, _, staff, password = setup(api, db, ["zone.add", "zone.edit", "zone.delete", "zone.export", "zone.import"])
    staff_headers = bearer(staff_login(api, staff, password))
    for denied in ({"Authorization": "Bearer " + mr_token}, staff_headers):
        for suffix in ("", "/sample", "/export", "/references/categories", "/references/locations"):
            assert api.get(BASE + suffix, headers=denied).status_code == 403
        for suffix in ("", "/import/review", "/import/commit"):
            assert api.post(BASE + suffix, headers=denied, json={}).status_code == 403
        identity = identity_from_token(db, denied["Authorization"].split()[1], settings)
        with pytest.raises(service.AllergenError) as cause:
            service.listing(db, identity)
        assert cause.value.status == 403


def test_download_acceptance_failure_releases_no_file(client, monkeypatch):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    from app.api.v1 import allergens as routes
    from fastapi import HTTPException
    def denied(*args, **kwargs):
        raise HTTPException(503, "Synthetic ledger unavailable")
    monkeypatch.setattr(routes, "server_record", denied)
    for suffix in ("/sample", "/export"):
        for fmt in ("csv", "xlsx"):
            result = api.get(BASE + suffix, headers=headers, params={"format": fmt})
            assert result.status_code == 503 and "Content-Disposition" not in result.headers
            assert result.headers["Content-Type"].startswith("application/json")


@pytest.mark.parametrize("file,filename", [
    (b"not csv", "file.csv"), (b"\xff", "file.csv"), (b"x" * (2 * 1024 * 1024 + 1), "file.csv"),
    (b"pretend excel", "file.xlsx"), (b"file", "file.xls"),
])
def test_invalid_and_oversized_files(client, file, filename):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    assert review(api, headers, file, filename).status_code in (413, 422)


def test_row_limit_duplicates_blank_money_and_inert_workbooks(client):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    c, l = references(api, headers)
    row = ["Test", c["name"], "", "0", l["name"], "1:10", "", "active", "No Mix"]
    for count, expected in ((1000, 200), (1001, 422)):
        response = review(api, headers, transfer.encode([transfer.HEADERS[:9]] + [row] * count, "csv"))
        assert response.status_code == expected
        if expected == 200:
            report = response.json()
            assert not report["valid"] and report["rows"][0]["selling_price"] is None
    for fmt in ("csv", "xlsx"):
        malformed = transfer.encode([transfer.HEADERS[:9], ["Bad", c["name"], "0.0000001", "101", l["name"], "", "-1", "active", "Allergens"]], fmt)
        assert not review(api, headers, malformed, f"bad.{fmt}").json()["valid"]
    book = load_workbook(io.BytesIO(transfer.sample("xlsx")))
    book.active["A2"] = "=HYPERLINK(\"https://invalid.example\", \"x\")"
    output = io.BytesIO()
    book.save(output); book.close()
    assert review(api, headers, output.getvalue(), "formula.xlsx").status_code == 422
    assert db.scalar(select(func.count()).select_from(AllergenProduct)) == 0


def test_original_numeric_workbook_lexemes_and_legacy_apostrophes(client):
    import zipfile
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    c, l = references(api, headers)
    text = transfer.encode([transfer.HEADERS[:9], [
        "Exact numeric", c["name"], "0", "0", l["name"], "1:10", "0", "active", "Mix"]], "xlsx")
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(text)) as source, zipfile.ZipFile(output, "w") as target:
        for entry in source.infolist():
            payload = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                from xml.etree import ElementTree
                root = ElementTree.fromstring(payload)
                for cell in root.iter():
                    coordinate = cell.attrib.get("r")
                    lexemes = {"C2": "9.99999999999999999E11", "D2": "9.9123456E1", "G2": "1E-6"}
                    if coordinate in lexemes:
                        cell.attrib["t"] = "n"
                        for child in list(cell):
                            cell.remove(child)
                        ElementTree.SubElement(cell, "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v").text = lexemes[coordinate]
                payload = ElementTree.tostring(root)
            target.writestr(entry, payload)
    parsed = transfer.parse(output.getvalue(), "numeric.xlsx")[0]
    assert parsed["selling_price"] == "999999999999.999999" and parsed["gst"] == "99.123456" and parsed["threshold_limit"] == "0.000001"
    report = review(api, headers, output.getvalue(), "numeric.xlsx").json()
    assert report["valid"]
    assert commit(api, headers, output.getvalue(), report["digest"], "numeric.xlsx").status_code == 200
    legacy = data(c, l, name="''Literal legacy", legacy=True)
    assert transfer.parse(legacy, "legacy.csv")[0]["name"] == "''Literal legacy"


def test_review_is_session_bound_and_multipart_is_supported(client):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    c, l = references(api, headers)
    file = data(c, l)
    response = api.post(BASE + "/import/review", headers=headers, params={"filename": "batch.csv"},
                        files={"file": ("batch.csv", file, "text/csv")})
    assert response.status_code == 200 and response.json()["valid"]
    previous = response.json()["digest"]
    from app.db.models import User
    actor = db.scalar(select(User).where(User.is_protected_system_admin.is_(True)))
    refreshed = login(api, actor.email)
    assert refreshed.status_code == 200
    new_headers = bearer(refreshed)
    assert commit(api, new_headers, file, previous, "batch.csv").status_code == 409
    assert db.scalar(select(func.count()).select_from(AllergenProduct)) == 0


def test_export_match_bound_does_not_log_failed_release_and_filter_can_narrow(client):
    import uuid
    from app.db.models import User
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    c, l = references(api, headers)
    actor = db.scalar(select(User).where(User.is_protected_system_admin.is_(True)))
    db.execute(AllergenProduct.__table__.insert(), [
        dict(id=uuid.uuid4(), name=f"Bound product {index}", category_id=c["id"], storage_location_id=l["id"],
             selling_price=None, gst=Decimal(0), concentration="1:10", threshold_limit=None, mix=False,
             status="active", version=1, created_by=actor.id, updated_by=actor.id)
        for index in range(5001)])
    # Publish synthetic rows to the enclosing fixture transaction; an expected
    # failed export rolls back its own savepoint, not this seeded catalogue.
    db.commit()
    for fmt in ("csv", "xlsx"):
        failure = api.get(BASE + "/export", headers=headers, params={"format": fmt})
        assert failure.status_code == 422 and failure.json()["error"]["code"] == "allergen_export_limit"
    assert db.scalar(select(func.count()).select_from(DownloadLog).where(DownloadLog.source == "allergen")) == 0
    page = api.get(BASE, headers=headers, params={"query": "Bound product 5000", "limit": 1}).json()
    assert page["total"] == 5001 and page["filtered"] == 1
    exported = api.get(BASE + "/export", headers=headers, params={"query": "Bound product 5000", "format": "csv"})
    assert exported.status_code == 200 and exported.headers["X-Download-Log"]
    assert len(transfer.parse(exported.content, "backup.csv")) == 1
    db.scalar(select(AllergenProduct).where(AllergenProduct.name == "Bound product 5000")).status = "inactive"
    db.commit()
    accepted = api.get(BASE + "/export", headers=headers, params={"status": "active", "format": "csv"})
    assert accepted.status_code == 200 and accepted.headers["X-Download-Log"]
    assert len(list(csv.reader(io.StringIO(accepted.content.decode("utf-8-sig"))))) == 5001
