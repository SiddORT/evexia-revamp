"""Synthetic isolated Vendor APIs; never configured database/account credentials."""
import csv
import io
import uuid
import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from app.db.models import AuditEvent, AuthSession
from app.db.vendor_models import Vendor
from app.db.download_models import DownloadLog
from app.schemas.vendors import VendorFields
from app.services import vendors, vendor_transfer
from test_sessions import client, create_user, login
from test_reporting import admin_headers
from test_zones import workbook, replace_zip

BASE = "/api/v1/admin/vendors"
FIELDS = dict(vendorName="Synthetic Supply", gstNo="27DDDDD3333D1Z8",
              registeredAddress="Block A\nMain Road", contactPersonName="Synthetic Contact",
              emailId="contact@example.test", phoneNo="+91 98765-43210", dialCountry="IN", status="active")
HEADER = ",".join(vendor_transfer.HEADERS[:8]) + "\n"


def add(api, headers, **changes):
    response = api.post(BASE, headers=headers, json={**FIELDS, **changes})
    assert response.status_code == 201, response.text
    return response.json()


def review(api, headers, data, filename="vendors.csv"):
    return api.post(BASE + "/import/review", headers={**headers, "Content-Type": "application/octet-stream"},
                    params={"filename": filename}, content=data)


def commit(api, headers, data, digest, filename="vendors.csv", **params):
    return api.post(BASE + "/import/commit", headers={**headers, "Content-Type": "application/octet-stream"},
                    params={"filename": filename, "digest": digest, "confirm": "true", **params}, content=data)


def test_crud_audit_literal_filters_pagination_and_soft_deletion(client):
    api, db, _ = client
    headers, actor = admin_headers(api, db)
    row = add(api, headers, vendorName="  City \t Supply ", gstNo="27ddddd3333d1z8")
    assert row["vendorName"] == "City Supply" and row["gstNo"] == FIELDS["gstNo"]
    assert row["phoneNo"] == "9876543210" and row["dialCountry"] == "IN"
    assert row["createdBy"] == row["updatedBy"] == "Super Admin"
    assert row["createdAt"] == row["updatedAt"] and row["status"] == "active"
    duplicate = api.post(BASE, headers=headers, json={**FIELDS, "vendorName": "CITY   supply", "gstNo": "29EEEEE4444E1Z9"})
    assert duplicate.status_code == 409
    assert api.post(BASE, headers=headers, json={**FIELDS, "vendorName": "Other"}).status_code == 409
    body = {**FIELDS, "vendorName": "Revised", "dialCountry": "GB", "phoneNo": "(7700) 900-123", "expected_version": 1}
    edited = api.post(f"{BASE}/{row['id']}/edit", headers=headers, json=body).json()
    assert edited["version"] == 2 and edited["phoneNo"] == "7700900123" and edited["dialCountry"] == "GB"
    assert edited["createdAt"] == row["createdAt"] and edited["updatedAt"] > row["updatedAt"]
    for query in ("revised", "DDDDD", "Main Road", "synthetic contact", "contact@example", "7700"):
        assert api.get(BASE, headers=headers, params={"query": query}).json()["filtered"] == 1
    for query in ("%", "_", "\\"):
        assert api.get(BASE, headers=headers, params={"query": query}).json()["filtered"] == 0
    for operation, data in (("edit", body), ("status", {"status": "inactive", "expected_version": 1}),
                            ("delete", {"expected_version": 1})):
        response = api.post(f"{BASE}/{row['id']}/{operation}", headers=headers, json=data)
        assert response.status_code == 409 and response.json()["error"]["code"] == "vendor_stale"
    assert api.post(f"{BASE}/{row['id']}/status", headers=headers,
                    json={"status": "inactive", "expected_version": 2}).json()["version"] == 3
    add(api, headers, vendorName="East", gstNo="29EEEEE4444E1Z9")
    add(api, headers, vendorName="West", gstNo="07FFFFF5555F1ZA")
    result = api.get(BASE, headers=headers, params={"status": "inactive", "query": "revised", "limit": 1}).json()
    assert result["filtered"] == 1 and result["total"] == 3
    pages = [api.get(BASE, headers=headers, params={"limit": 2, "offset": offset}).json() for offset in (0, 2)]
    assert len({r["id"] for page in pages for r in page["items"]}) == 3
    assert api.post(f"{BASE}/{row['id']}/delete", headers=headers, json={"expected_version": 3}).json()["version"] == 4
    assert api.get(f"{BASE}/{row['id']}", headers=headers).status_code == 404
    assert api.get(BASE, headers=headers).json()["total"] == 2
    assert "Revised" not in api.get(BASE + "/export", headers=headers).text
    db.expire_all()
    stored = db.get(Vendor, uuid.UUID(row["id"]))
    assert stored.deleted_at == stored.updated_at and stored.deleted_by == actor.id
    event = db.scalar(select(AuditEvent).where(AuditEvent.action == "vendor_delete"))
    assert event.actor_id == actor.id and event.session_id and event.resource_type == "vendor"
    assert add(api, headers, vendorName="Revised")["id"] != row["id"]


@pytest.mark.parametrize("country,phone,expected", [
    ("IN", "+91 (98765) 43210", "9876543210"), ("US", "(202) 555-0123", "2025550123"),
    ("GB", "7700 900123", "7700900123"), ("AE", "50 123 4567", "501234567")])
def test_all_phone_countries_round_trip(client, country, phone, expected):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    row = add(api, headers, dialCountry=country, phoneNo=phone)
    assert row["dialCountry"] == country and row["phoneNo"] == expected
    for format in ("csv", "xlsx"):
        exported = api.get(BASE + "/export", headers=headers, params={"format": format})
        parsed = vendor_transfer.parse(exported.content, "vendors." + format)
        assert parsed[0]["values"]["dialCountry"] == country and parsed[0]["values"]["phoneNo"] == expected
        assert parsed[0]["errors"] == []


def test_required_bounds_phone_errors_and_forged_fields(client):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    for field in vendor_transfer.LABELS:
        if field in ("dialCountry", "status"):
            continue
        assert api.post(BASE, headers=headers, json={**FIELDS, field: ""}).status_code == 422
    for field, value in (("vendorName", "x"*201), ("registeredAddress", "x"*2001), ("contactPersonName", "x"*201),
                         ("emailId", "x"*321), ("emailId", "bad"), ("gstNo", "invalid"), ("status", "wrong"),
                         ("dialCountry", "ZZ"), ("phoneNo", "123"), ("phoneNo", "1234567890"),
                         ("phoneNo", "+1 2025550123"), ("vendorName", "nul\0text"),
                         ("registeredAddress", "bad\x01text")):
        response = api.post(BASE, headers=headers, json={**FIELDS, field: value})
        assert response.status_code == 422, (field, response.text)
        assert "input" not in response.json()["error"]
    for key in ("createdBy", "updatedAt", "created_by", "deleted_by", "id", "version", "password_secret"):
        response = api.post(BASE, headers=headers, json={**FIELDS, key: "private-value"})
        assert response.status_code == 422 and "private-value" not in response.text and "password_secret" not in response.text
    # Optional defaults are explicit, never an unknown country fallback.
    fields = {k: v for k, v in FIELDS.items() if k not in ("dialCountry", "status")}
    assert api.post(BASE, headers=headers, json=fields).json()["status"] == "active"


def test_anonymous_mr_non_singleton_and_revoked_service_denied(client):
    api, db, settings = client
    assert api.get(BASE).status_code == 401
    user = create_user(db, "vendor-mr@example.com")
    denied = {"Authorization": "Bearer " + login(api, user.email).json()["access_token"]}
    for route in ("", "/sample", "/export", "/" + str(uuid.uuid4())):
        assert api.get(BASE + route, headers=denied).status_code == 403
    assert api.post(BASE, headers=denied, json=FIELDS).status_code == 403
    assert review(api, denied, b"garbage").status_code == 403
    from app.services.auth import identity_from_token
    actor = identity_from_token(db, denied["Authorization"][7:], settings)
    for work in (lambda: vendors.create(db, actor, VendorFields(**FIELDS)),
                 lambda: vendor_transfer.export(db, actor, "", "all", "csv"),
                 lambda: vendor_transfer.transfer(db, actor, vendor_transfer.sample("csv"), "vendors.csv")):
        with pytest.raises(vendors.VendorError) as failure:
            work()
        assert failure.value.status == 403
    headers, _ = admin_headers(api, db)
    actor = identity_from_token(db, headers["Authorization"][7:], settings)
    session = db.get(AuthSession, actor.session_id)
    from app.core.security import utcnow
    session.status, session.revoked_at = "REVOKED", utcnow()
    db.commit()
    from app.services.auth import AuthError
    with pytest.raises(AuthError):
        vendors.create(db, actor, VendorFields(**FIELDS))


def test_inert_atomic_review_file_filename_session_and_duplicates(client):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    data = (HEADER + "First,27DDDDD3333D1Z8,Address,Contact,one@example.test,9876543210,IN,Active\n"
                    "Second,29EEEEE4444E1Z9,Address,Contact,two@example.test,501234567,AE,Inactive").encode()
    report = review(api, headers, data).json()
    assert report["valid"] and report["validCount"] == 2 and report["invalidCount"] == 0
    assert db.scalar(select(func.count()).select_from(Vendor)) == 0
    assert commit(api, headers, data + b"\n", report["digest"]).status_code == 409
    assert commit(api, headers, data, report["digest"], filename="changed.csv").status_code == 409
    assert commit(api, headers, data, report["digest"], confirm="false").status_code == 422
    add(api, headers, vendorName="Second", gstNo="07FFFFF5555F1ZA")
    assert commit(api, headers, data, report["digest"]).status_code == 409
    assert not db.scalar(select(Vendor.id).where(Vendor.vendorName == "First"))
    duplicate = (HEADER + " Same,27DDDDD3333D1Z8,A,C,a@example.test,9876543210,IN,active\n"
                           "SAME,27ddddd3333d1z8,A,C,b@example.test,9876543210,IN,inactive").encode()
    report = review(api, headers, duplicate).json()
    assert report["invalidCount"] == 1 and any("GST" in e for e in report["rows"][1]["errors"])
    good = review(api, headers, vendor_transfer.sample("csv")).json()
    replacement, _ = admin_headers(api, db)
    assert commit(api, replacement, vendor_transfer.sample("csv"), good["digest"]).status_code == 409


def test_legacy_current_csv_xlsx_round_trip_and_historical_audit_ignored(client):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    legacy = (",".join(vendor_transfer.HEADERS[:6]) + "\nLegacy,27DDDDD3333D1Z8,A,C,c@example.test,+91 9876543210").encode()
    parsed = review(api, headers, legacy).json()
    assert parsed["valid"] and parsed["rows"][0]["values"]["dialCountry"] == "IN"
    assert commit(api, headers, legacy, parsed["digest"]).json()["imported"] == 1
    for format in ("csv", "xlsx"):
        sample = api.get(BASE + "/sample", headers=headers, params={"format": format})
        assert sample.status_code == 200 and sample.headers.get("X-Download-Log")
        report = review(api, headers, sample.content, filename="sample." + format).json()
        assert report["valid"]
        exported = api.get(BASE + "/export", headers=headers, params={"format": format})
        assert exported.status_code == 200 and exported.headers.get("X-Download-Log")
        if format == "xlsx":
            book = load_workbook(io.BytesIO(exported.content))
            assert all(c.data_type == "s" for row in book.active for c in row)
            book.close()
        assert "session" not in exported.text if format == "csv" else True
        assert vendor_transfer.parse(exported.content, "export." + format)[0]["errors"] == []
    # Formula-like ordinary text remains inert and round-trips, including apostrophes.
    for name in ("=Supply", "'Supply", "@Supply"):
        row = {**FIELDS, "vendorName": name, "registeredAddress": "'=literal\naddress"}
        values = [row[field] for field in vendor_transfer.BUSINESS_FIELDS]
        for format in ("csv", "xlsx"):
            data = vendor_transfer.encode([vendor_transfer.HEADERS,
                [vendor_transfer.safe_text(v) for v in values] + ["Forged actor", "1900", "Forged actor", "1900"]], format)
            result = vendor_transfer.parse(data, "transfer." + format)[0]
            assert result["values"]["vendorName"] == name and result["values"]["registeredAddress"] == row["registeredAddress"]
    current = vendor_transfer.encode([vendor_transfer.HEADERS,
        ["Current", "29EEEEE4444E1Z9", "A", "C", "c@example.test", "2025550123", "US", "inactive",
         "Forged actor", "1900", "Forged actor", "1900"]], "csv")
    report = review(api, headers, current).json()
    assert commit(api, headers, current, report["digest"]).json()["imported"] == 1
    created = api.get(BASE, headers=headers, params={"query": "Current"}).json()["items"][0]
    assert created["createdBy"] == "Super Admin" and not created["createdAt"].startswith("1900")


def test_upload_attacks_limits_and_download_fail_closed(client, monkeypatch):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    for data, filename, status in ((b"", "vendors.csv", 413), (b"x"*(2*1024*1024+1), "vendors.csv", 413),
                                   (b"bad", "vendors.xls", 422), (b"bad", "vendors.xlsx", 422),
                                   (b"\xff", "vendors.csv", 422), (b'Bad,Headers\n"x', "vendors.csv", 422),
                                   ((HEADER + ("A,B,C,D,E,F,G,H\n"*1001)).encode(), "vendors.csv", 422)):
        assert review(api, headers, data, filename).status_code == status
    bad = (HEADER + "A,27DDDDD3333D1Z8,A,C,invalid,123,ZZ,active").encode()
    report = review(api, headers, bad).json()
    assert not report["valid"] and report["invalidCount"] == 1
    xlsx = workbook([vendor_transfer.HEADERS[:8], ["A", "27DDDDD3333D1Z8", "=1+1", "C", "c@example.test", "9876543210", "IN", "active"]])
    assert review(api, headers, xlsx, "attack.xlsx").status_code == 422
    for files in ({"file": ("sample.csv", vendor_transfer.sample("csv"), "text/csv"), "other": ("x", b"x")},
                  {"wrong": ("sample.csv", vendor_transfer.sample("csv"), "text/csv")}):
        assert api.post(BASE + "/import/review", headers=headers, params={"filename": "sample.csv"}, files=files).status_code == 422
    response = api.post(BASE + "/import/review", headers=headers, params={"filename": "sample.csv"},
                        files={"file": ("sample.csv", vendor_transfer.sample("csv"), "text/csv")})
    assert response.status_code == 200 and response.json()["valid"]
    from app.api.v1 import vendors as routes
    from fastapi import HTTPException
    def denied(*args, **kwargs):
        raise HTTPException(503, "Download acceptance unavailable")
    monkeypatch.setattr(routes, "server_record", denied)
    for path in ("/sample", "/export"):
        assert api.get(BASE + path, headers=headers).status_code == 503


def test_complete_filtered_exports_limit_and_evidence_dedup(client, monkeypatch):
    api, db, _ = client
    headers, _ = admin_headers(api, db)
    for index in range(12):
        add(api, headers, vendorName=f"Matching {index}", gstNo=f"27DDDDD{index:04d}D1Z8",
            status="inactive" if index < 11 else "active")
    assert len(api.get(BASE, headers=headers, params={"limit": 2}).json()["items"]) == 2
    initiation = str(uuid.uuid4())
    for format in ("csv", "xlsx"):
        response = api.get(BASE + "/export", headers={**headers, "X-Download-Initiation": initiation},
                           params={"query": "Matching", "status": "inactive", "format": format})
        assert response.status_code == 200
        if format == "csv":
            assert len(list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))) == 12
            repeated = api.get(BASE + "/export", headers={**headers, "X-Download-Initiation": initiation},
                               params={"query": "Matching", "status": "inactive", "format": format})
            assert repeated.headers["X-Download-Log"] == response.headers["X-Download-Log"]
        else:
            book = load_workbook(io.BytesIO(response.content))
            assert book.active.max_row == 12
            book.close()
        initiation = str(uuid.uuid4())
    assert db.scalar(select(func.count()).select_from(DownloadLog).where(DownloadLog.source == "vendor")) == 2
    monkeypatch.setattr(vendor_transfer, "EXPORT_LIMIT", 10)
    assert api.get(BASE + "/export", headers=headers, params={"status": "inactive"}).status_code == 422


def test_staff_all_catalogue_grants_cannot_access_vendor(client):
    from test_zone_permissions import setup, staff_login, bearer
    from app.core.master_catalogue import MASTER_ACTIONS
    api, db, _ = client
    admin, _, _, record, password = setup(api, db, list(MASTER_ACTIONS))
    row = add(api, admin)
    staff = bearer(staff_login(api, record, password))
    for route in ("", "/" + row["id"], "/sample", "/export"):
        assert api.get(BASE + route, headers=staff).status_code == 403
    for route, body in (("", FIELDS), ("/" + row["id"] + "/edit", {**FIELDS, "expected_version": 1}),
                        ("/" + row["id"] + "/status", {"status": "inactive", "expected_version": 1}),
                        ("/" + row["id"] + "/delete", {"expected_version": 1})):
        assert api.post(BASE + route, headers=staff, json=body).status_code == 403
    assert review(api, staff, vendor_transfer.sample("csv")).status_code == 403
    assert commit(api, staff, vendor_transfer.sample("csv"), "0"*64).status_code == 403
