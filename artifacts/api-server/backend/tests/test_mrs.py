"""Synthetic MR account/directory API and parser regressions, isolated PostgreSQL."""
import csv
import io
import json
import uuid
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import func, select
from app.db.models import User, MRProfile, Patient, AuditEvent
from app.db.mr_models import MRDirectory
from app.schemas.mrs import MRFields
from app.services import mr_transfer, mr_postal, mrs
from app.core.security import verify_password
from test_sessions import client, PASSWORD, create_user
from test_reporting import admin_headers
from test_zones import workbook, replace_zip

BASE = "/api/v1/admin/mrs"


def setup(api, db):
    headers, actor = admin_headers(api, db)
    hq = api.post("/api/v1/admin/headquarters", headers=headers, json={"name": "HQ for MRs", "status": "active"}).json()["id"]
    zone = api.post("/api/v1/admin/zones", headers=headers, json={"name": "Zone for MRs", "status": "active"}).json()["id"]
    return headers, actor, hq, zone


def fields(hq, zone, name="Synthetic MR", code="MR-01", username="synthetic.mr", **changes):
    return dict(name=name, employeeCode=code, userId=username, phone="", email="", contactRequirement="optional",
                hq=hq, zoneId=zone, dateOfJoining="2020-01-01", designation="Medical Representative",
                reportingManagerId=None, paymentLimit="", doctorDaysLimit="", status="active",
                pincode="110001", addressLine1="Synthetic address", addressLine2="", landmark="Synthetic landmark",
                city="Delhi", state="Delhi", country="India", **changes)


def add(api, headers, body):
    response = api.post(BASE, headers=headers, json=body)
    assert response.status_code == 201, response.text
    assert response.headers["cache-control"] == "no-store"
    return response.json()


def mr_login(api, identifier, password, portal="mr"):
    return api.post("/api/v1/auth/login", headers={"Origin": "http://testserver"},
                    json={"identifier": identifier, "password": password, "identity_kind": portal})


def review(api, headers, data, filename="mrs.csv"):
    return api.post(BASE + "/import/review", headers={**headers, "Content-Type": "application/octet-stream"},
                    content=data, params={"filename": filename})


def commit(api, headers, data, digest, filename="mrs.csv"):
    return api.post(BASE + "/import/commit", headers={**headers, "Content-Type": "application/octet-stream"},
                    content=data, params={"filename": filename, "digest": digest, "confirm": "true"})


def csv_data(bodies, headers=None):
    headers = headers or mr_transfer.HEADERS
    stream = io.StringIO()
    writer = csv.writer(stream)
    writer.writerow(headers)
    for body in bodies:
        values = dict(body, hq="HQ for MRs", zoneId="Zone for MRs",
                      reportingManagerId=body["reportingManagerId"] or "")
        writer.writerow([values[key] for key, title in mr_transfer.COLUMNS if title in headers])
    return stream.getvalue().encode()


def test_persisted_directory_account_identity_and_secret_projection(client):
    api, db, _ = client
    headers, actor, hq, zone = setup(api, db)
    legacy = create_user(db, "existing-identity-only@example.invalid")
    patient = Patient(assigned_mr_id=db.scalar(select(MRProfile.id).where(MRProfile.user_id == legacy.id)), is_active=True)
    db.add(patient)
    db.commit()
    existing_profile = patient.assigned_mr_id
    result = add(api, headers, {**fields(hq, zone), "initialPassword": "  spaces retained in password  "})
    row, secret = result["record"], result["credentials"]
    assert secret["password"] == "  spaces retained in password  "
    assert row["paymentLimit"] == "0.00" and row["doctorDaysLimit"] == 0
    assert row["createdBy"] == row["updatedBy"] == "Super Admin"
    db.expire_all()
    persisted = db.get(MRDirectory, uuid.UUID(row["id"]))
    profile = db.get(MRProfile, persisted.id)
    user = db.get(User, profile.user_id)
    assert user.email is None and user.username == "synthetic.mr" and user.system_role == "mr"
    assert verify_password(secret["password"], user.password_hash)
    assert db.get(MRDirectory, existing_profile) is None
    assert db.get(Patient, patient.id).assigned_mr_id == existing_profile
    for path in (BASE, BASE + "/" + row["id"]):
        response = api.get(path, headers=headers)
        assert response.status_code == 200
        assert secret["password"] not in response.text and "password_hash" not in response.text
    logged = mr_login(api, secret["userId"], secret["password"])
    assert logged.status_code == 200, logged.text
    assert logged.json()["user"]["mr_id"] == row["id"]
    assert logged.json()["user"]["email"] is None
    mr_headers = {"Authorization": "Bearer " + logged.json()["access_token"]}
    assert api.get(BASE, headers=mr_headers).status_code == 403
    assert api.get("/api/v1/admin/headquarters", headers=mr_headers).status_code == 403
    assert mr_login(api, secret["userId"], secret["password"], "admin").status_code == 401
    assert mr_login(api, actor.email, PASSWORD, "mr").status_code == 401
    assert api.get("/api/v1/auth/me", headers=mr_headers).status_code == 200
    changed = api.post("/api/v1/auth/change-password", headers=mr_headers,
                       json={"current_password": secret["password"], "new_password": "Synthetic changed password"})
    assert changed.status_code == 204
    assert api.get("/api/v1/auth/me", headers=mr_headers).status_code == 401


@pytest.mark.parametrize("field,value", [
    ("name", ""), ("name", "x" * 201), ("name", "a\x00b"), ("employeeCode", "x" * 65),
    ("userId", "email@example.invalid"), ("userId", "UPPER?!"), ("email", "bad"), ("phone", "123"),
    ("pincode", "012345"), ("pincode", "12345"), ("dateOfJoining", "2020-02-30"),
    ("dateOfJoining", "2020-01-01T12:00:00Z"), ("dateOfJoining", (date.today() + timedelta(days=3)).isoformat()),
    ("paymentLimit", "0.001"), ("paymentLimit", "-1"), ("paymentLimit", "1e2"),
    ("paymentLimit", "1000000000.00"), ("doctorDaysLimit", 3651), ("doctorDaysLimit", 0.5),
    ("doctorDaysLimit", True), ("country", "x" * 101), ("landmark", ""),
    ("initialPassword", "short"), ("arbitrary_user_id", str(uuid.uuid4())),
])
def test_business_validation(client, field, value):
    api, db, _ = client
    headers, _, hq, zone = setup(api, db)
    response = api.post(BASE, headers=headers, json={**fields(hq, zone), field: value})
    assert response.status_code == 422, response.text
    assert db.scalar(select(func.count()).select_from(MRDirectory)) == 0


def test_optional_contacts_normalization_and_uniqueness(client):
    api, db, _ = client
    headers, _, hq, zone = setup(api, db)
    body = fields(hq, zone)
    required = api.post(BASE, headers=headers, json={**body, "contactRequirement": "required"})
    assert required.status_code == 422
    first = add(api, headers, {**body, "phone": "+91 (98765) 43210", "email": " Test@Example.COM ",
                               "name": "  Same name  ", "employeeCode": " Mr-01 "})["record"]
    assert first["phone"] == "9876543210" and first["email"] == "test@example.com"
    for changes in ({"employeeCode": "mr-01", "userId": "different.mr", "email": ""},
                    {"employeeCode": "MR-02", "userId": "SYNTHETIC.MR", "email": ""},
                    {"employeeCode": "MR-02", "userId": "different.mr", "email": "TEST@EXAMPLE.COM"},
                    {"employeeCode": "MR-02", "userId": "root", "email": ""}):
        assert api.post(BASE, headers=headers, json={**body, **changes}).status_code == 409
    add(api, headers, fields(hq, zone, name="Same name", code="MR-02", username="second.mr"))
    assert api.get(BASE, headers=headers).json()["total"] == 2
    assert "password" not in api.get(BASE + "/username", headers=headers).json()


def test_versions_assignment_lifecycle_cycles_and_delete_blocks(client):
    api, db, _ = client
    headers, _, hq, zone = setup(api, db)
    a = add(api, headers, fields(hq, zone))["record"]
    b = add(api, headers, {**fields(hq, zone, code="MR-02", username="second.mr"), "reportingManagerId": a["id"]})["record"]
    cyclic = {**fields(hq, zone), "expected_version": a["version"], "reportingManagerId": b["id"]}
    assert api.post(BASE + "/" + a["id"] + "/edit", headers=headers, json=cyclic).status_code == 409
    assert api.post(BASE + "/" + a["id"] + "/delete", headers=headers, json={"expected_version": 1}).status_code == 409
    assert api.post(BASE + "/" + a["id"] + "/contact", headers=headers,
                    json={"expected_version": 1, "contactRequirement": "required"}).status_code == 422
    assert api.post("/api/v1/admin/headquarters/" + hq + "/status", headers=headers,
                    json={"expected_version": 1, "status": "inactive"}).status_code == 200
    retained = api.post(BASE + "/" + a["id"] + "/edit", headers=headers,
                        json={**fields(hq, zone), "expected_version": 1})
    assert retained.status_code == 200 and retained.json()["assignmentWarnings"]
    assert api.post(BASE + "/" + a["id"] + "/status", headers=headers,
                    json={"expected_version": 1, "status": "inactive"}).status_code == 409
    assert api.post(BASE, headers=headers, json=fields(hq, zone, code="MR-03", username="third.mr")).status_code == 409
    assert api.post("/api/v1/admin/headquarters/" + hq + "/delete", headers=headers,
                    json={"expected_version": 2}).status_code == 200
    assert api.post(BASE + "/" + a["id"] + "/edit", headers=headers,
                    json={**fields(hq, zone), "expected_version": 2}).status_code == 409
    choices = api.get(BASE + "/references", headers=headers, params={"kind": "headquarters", "include_saved": hq}).json()
    assert choices["items"][0]["deleted"]
    assert api.post(BASE + "/" + b["id"] + "/delete", headers=headers, json={"expected_version": 1}).status_code == 200
    deleted = api.post(BASE + "/" + a["id"] + "/delete", headers=headers, json={"expected_version": 2})
    assert deleted.status_code == 200
    db.expire_all()
    row = db.get(MRDirectory, uuid.UUID(a["id"]))
    assert row.deleted_at == row.updated_at and row.deleted_by == row.updated_by
    profile = db.get(MRProfile, row.id)
    assert not profile.is_active and not db.get(User, profile.user_id).is_active
    assert api.get(BASE + "/" + a["id"], headers=headers).status_code == 404
    assert api.post(BASE + "/" + a["id"] + "/status", headers=headers,
                    json={"expected_version": row.version, "status": "active"}).status_code == 404


def test_reset_deactivation_identifier_changes_revoke_sessions(client):
    api, db, _ = client
    headers, _, hq, zone = setup(api, db)
    result = add(api, headers, fields(hq, zone))
    row, secret = result["record"], result["credentials"]
    token = mr_login(api, secret["userId"], secret["password"]).json()["access_token"]
    mh = {"Authorization": "Bearer " + token}
    reset = api.post(BASE + "/" + row["id"] + "/reset", headers=headers, json={"expected_version": 1})
    assert reset.status_code == 200 and reset.json()["credentials"]["password"] != secret["password"]
    assert api.get("/api/v1/auth/me", headers=mh).status_code == 401
    new_secret = reset.json()["credentials"]
    token = mr_login(api, new_secret["userId"], new_secret["password"]).json()["access_token"]
    edit = api.post(BASE + "/" + row["id"] + "/edit", headers=headers,
                    json={**fields(hq, zone, username="changed.mr"), "expected_version": 2})
    assert edit.status_code == 200 and edit.json()["id"] == row["id"] and "credentials" not in edit.json()
    assert api.get("/api/v1/auth/me", headers={"Authorization": "Bearer " + token}).status_code == 401
    assert mr_login(api, "changed.mr", new_secret["password"]).status_code == 200
    inactive = api.post(BASE + "/" + row["id"] + "/status", headers=headers, json={"expected_version": 3, "status": "inactive"})
    assert inactive.status_code == 200
    assert api.post("/api/v1/auth/refresh", headers={"Origin": "http://testserver"}).status_code == 401
    assert mr_login(api, "changed.mr", new_secret["password"]).status_code == 401
    active = api.post(BASE + "/" + row["id"] + "/status", headers=headers, json={"expected_version": 4, "status": "active"})
    assert active.status_code == 200 and mr_login(api, "changed.mr", new_secret["password"]).status_code == 200


def test_server_combined_filters_full_export_and_ledger(client):
    api, db, _ = client
    headers, _, hq, zone = setup(api, db)
    add(api, headers, fields(hq, zone))
    add(api, headers, {**fields(hq, zone, code="MR-02", username="second.mr"), "status": "inactive"})
    matching = api.get(BASE, headers=headers, params={"query": "synthetic", "status": "inactive",
                       "zone_id": zone, "hq_id": hq, "limit": 1}).json()
    assert matching["filtered"] == 1 and matching["total"] == 2 and len(matching["items"]) == 1
    assert api.get(BASE, headers=headers, params={"query": "%"}).json()["filtered"] == 0
    assert api.get(BASE, headers=headers, params={"query": "second.mr"}).json()["filtered"] == 1
    for format in ("csv", "xlsx"):
        response = api.get(BASE + "/export", headers={**headers, "X-Download-Initiation": str(uuid.uuid4())},
                           params={"format": format, "zone_id": zone, "hq_id": hq})
        assert response.status_code == 200 and response.headers["X-Download-Log"]
        if format == "csv":
            exported = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
            assert exported[0] == mr_transfer.HEADERS and len(exported) == 3
        else:
            book = load_workbook(io.BytesIO(response.content))
            assert book.active.max_row == 3 and book.active.max_column == 25
            book.close()
    for format in ("csv", "xlsx"):
        sample = api.get(BASE + "/sample", headers=headers, params={"format": format})
        assert sample.status_code == 200 and sample.headers["X-Download-Log"]
        assert len(mr_transfer.parse(sample.content, "sample." + format)) == 1


def test_atomic_import_legacy_exports_batch_managers_and_secret_isolation(client):
    api, db, _ = client
    headers, _, hq, zone = setup(api, db)
    one = fields(hq, zone)
    two = fields(hq, zone, code="MR-02", username="second.mr")
    two["reportingManagerId"] = "user:synthetic.mr"
    data = csv_data([one, two])
    checked = review(api, headers, data)
    assert checked.status_code == 200 and checked.json()["valid"], checked.text
    assert db.scalar(select(func.count()).select_from(MRDirectory)) == 0
    assert commit(api, headers, data + b"\n", checked.json()["digest"]).status_code == 409
    saved = commit(api, headers, data, checked.json()["digest"])
    assert saved.status_code == 200 and saved.json()["imported"] == 2, saved.text
    assert saved.headers["cache-control"] == "no-store"
    assert len(saved.json()["credentials"]) == 2
    assert commit(api, headers, data, checked.json()["digest"]).status_code == 409
    assert db.scalar(select(func.count()).select_from(MRDirectory)) == 2
    rows = list(db.scalars(select(MRDirectory).order_by(MRDirectory.employeeCode)))
    assert rows[1].reportingManagerId == rows[0].id
    exported = api.get(BASE + "/export", headers=headers).content
    assert [r["values"]["userId"] for r in mr_transfer.parse(exported, "mrs.csv")]
    legacy = fields(hq, zone, code="MR-03", username="legacy.mr")
    legacy.update(phone="9876543210", email="legacy@example.com")
    check = review(api, headers, csv_data([legacy], mr_transfer.LEGACY))
    assert check.status_code == 200 and check.json()["valid"], check.text
    assert mr_transfer.parse(csv_data([legacy], mr_transfer.LEGACY), "old.csv")[0]["values"]["contactRequirement"] == "required"
    audits = json.dumps([dict(action=e.action, reason=e.reason) for e in db.scalars(select(AuditEvent))])
    for credentials in saved.json()["credentials"]:
        assert credentials["password"] not in audits
        assert credentials["password"].encode() not in exported


def test_import_cycles_duplicates_invalid_and_changed_relationships(client):
    api, db, _ = client
    headers, _, hq, zone = setup(api, db)
    one, two = fields(hq, zone), fields(hq, zone, code="MR-02", username="second.mr")
    one["reportingManagerId"], two["reportingManagerId"] = "user:second.mr", "user:synthetic.mr"
    response = review(api, headers, csv_data([one, two]))
    assert response.status_code == 200 and not response.json()["valid"]
    assert all(row["errors"] for row in response.json()["rows"])
    assert not review(api, headers, csv_data([fields(hq, zone), fields(hq, zone)])).json()["valid"]
    data = csv_data([fields(hq, zone)])
    checked = review(api, headers, data).json()
    api.post("/api/v1/admin/headquarters/" + hq + "/status", headers=headers, json={"expected_version": 1, "status": "inactive"})
    assert commit(api, headers, data, checked["digest"]).status_code == 409
    assert db.scalar(select(func.count()).select_from(MRDirectory)) == 0


def test_transfer_boundaries_and_inert_formula_roundtrip():
    with pytest.raises(mrs.MRError):
        mr_transfer.parse(b"x" * (mr_transfer.MAX_BYTES + 1), "large.csv")
    with pytest.raises(mrs.MRError):
        mr_transfer.parse(b"MR Name,Email ID,Phone No.,HQ,Assigned Zone,Status\nx,x,x,x,x,x", "mock.csv")
    with pytest.raises(mrs.MRError):
        mr_transfer.parse((",".join(mr_transfer.HEADERS + ["Password"]) + "\n").encode(), "password.csv")
    data = mr_transfer.sample("xlsx")
    assert mr_transfer.parse(data, "real.xlsx")
    book = load_workbook(io.BytesIO(data))
    book.active["B2"] = "=WEBSERVICE(\"https://example.invalid\")"
    out = io.BytesIO()
    book.save(out)
    book.close()
    with pytest.raises(mrs.MRError):
        mr_transfer.parse(out.getvalue(), "formula.xlsx")
    values = [""] * 21
    values[1] = "=literal business label"
    parsed = mr_transfer.parse(mr_transfer.encode([values], "csv"), "safe.csv")
    assert parsed[0]["values"]["name"] == "=literal business label"


def test_postal_unambiguous_choices_manual_failure_and_revocation(client, monkeypatch):
    api, db, _ = client
    headers, _, _, _ = setup(api, db)
    mr_postal._cache.clear()
    calls = []
    def provider(pin):
        calls.append(pin)
        return [{"city": "District suggestion", "state": "State", "country": "India"}]
    monkeypatch.setattr(mr_postal, "provider", provider)
    # Default argument is bound at definition; inject through the route adapter.
    original = mr_postal.lookup
    monkeypatch.setattr(mr_postal, "lookup", lambda db, actor, pin: original(db, actor, pin, provider))
    for _ in range(2):
        response = api.get(BASE + "/postal/110001", headers=headers)
        assert response.status_code == 200 and response.json()["choices"][0]["city"] == "District suggestion"
    assert calls == ["110001"]
    assert api.get(BASE + "/postal/invalid", headers=headers).status_code == 422
    def failed(pin):
        raise TimeoutError()
    monkeypatch.setattr(mr_postal, "lookup", lambda db, actor, pin: original(db, actor, pin, failed))
    response = api.get(BASE + "/postal/110002", headers=headers)
    assert response.status_code == 200 and response.json()["choices"] == []
    assert "manually" in response.json()["message"]


def test_auth_anonymous_and_email_optional_disabled_create(client):
    api, db, _ = client
    assert api.get(BASE).status_code == 401
    headers, _, hq, zone = setup(api, db)
    row = add(api, headers, {**fields(hq, zone), "status": "inactive"})
    assert mr_login(api, row["credentials"]["userId"], row["credentials"]["password"]).status_code == 401
    user = create_user(db, "plain-mr@example.com")
    token = mr_login(api, user.username, PASSWORD).json()["access_token"]
    denied = {"Authorization": "Bearer " + token}
    for path in ("", "/references?kind=zones", "/username", "/sample", "/export", "/postal/110001"):
        assert api.get(BASE + path, headers=denied).status_code == 403


def test_import_request_review_session_export_and_rate_bounds(client, monkeypatch):
    api, db, _ = client
    headers, _, hq, zone = setup(api, db)
    data = csv_data([fields(hq, zone)])
    report = review(api, headers, data).json()
    new_headers, _ = admin_headers(api, db)
    assert commit(api, new_headers, data, report["digest"]).status_code == 409
    assert review(api, new_headers, b"x" * (mr_transfer.MAX_BYTES + 1)).status_code == 413
    assert db.scalar(select(func.count()).select_from(MRDirectory)) == 0
    add(api, new_headers, fields(hq, zone))
    add(api, new_headers, fields(hq, zone, code="MR-02", username="second.mr"))
    monkeypatch.setattr(mr_transfer, "EXPORT_LIMIT", 1)
    assert api.get(BASE + "/export", headers=new_headers).status_code == 409
    mr_postal._cache.clear()
    original = mr_postal.lookup
    monkeypatch.setattr(mr_postal, "lookup", lambda db, actor, pin: original(db, actor, pin, lambda _: []))
    for _ in range(30):
        assert api.get(BASE + "/postal/110001", headers=new_headers).status_code == 200
    assert api.get(BASE + "/postal/110001", headers=new_headers).status_code == 429


def test_actual_postal_fixed_provider_district_dedupe_and_payload_guards(monkeypatch):
    import httpx
    original = httpx.Client
    calls = []
    def transport(request):
        calls.append(request)
        return httpx.Response(200, stream=httpx.ByteStream(json.dumps([{"Status": "Success", "PostOffice": [
            {"Name": "Office one", "District": "District", "State": "State", "Country": "India"},
            {"Name": "Office two", "District": "District", "State": "State", "Country": "India"},
        ]}]).encode()), headers={"Content-Type": "application/json"})
    monkeypatch.setattr(mr_postal.httpx, "Client", lambda **options: original(**options, transport=httpx.MockTransport(transport)))
    assert mr_postal.provider("110001") == [{"city": "District", "state": "State", "country": "India"}]
    assert str(calls[0].url) == "https://api.postalpincode.in/pincode/110001"
    assert calls[0].headers["accept-encoding"] == "identity"
    assert calls[0].content == b""
    def unsafe(request):
        return httpx.Response(302, headers={"Location": "http://127.0.0.1/private"})
    monkeypatch.setattr(mr_postal.httpx, "Client", lambda **options: original(**options, transport=httpx.MockTransport(unsafe)))
    with pytest.raises(Exception):
        mr_postal.provider("110001")


def test_parser_record_cell_macro_external_and_audit_bounds():
    from zipfile import ZipFile, ZIP_DEFLATED
    too_many = (",".join(mr_transfer.HEADERS) + "\n" + (",".join([""] * 21) + "\n") * 1001).encode()
    with pytest.raises(mrs.MRError):
        mr_transfer.parse(too_many, "too-many.csv")
    cell = (",".join(mr_transfer.HEADERS) + "\n" + "x" * 10001 + "," * 20).encode()
    with pytest.raises(mrs.MRError):
        mr_transfer.parse(cell, "cell.csv")
    for member in ("xl/vbaProject.bin", "xl/externalLinks/externalLink1.xml"):
        source = mr_transfer.sample("xlsx")
        out = io.BytesIO()
        with ZipFile(io.BytesIO(source)) as old, ZipFile(out, "w", ZIP_DEFLATED) as new:
            for item in old.infolist():
                new.writestr(item.filename, old.read(item))
            new.writestr(member, b"<unsafe/>")
        with pytest.raises(mrs.MRError):
            mr_transfer.parse(out.getvalue(), "unsafe.xlsx")
    parsed = mr_transfer.parse(mr_transfer.sample("xlsx"), "audit.xlsx")
    assert "createdBy" not in parsed[0]["values"]


def test_xlsx_payment_uses_exact_xml_not_rounded_parser_floats():
    data = workbook([mr_transfer.HEADERS, [""] * 11 + [1.0] + [""] * 9])
    from defusedxml import ElementTree
    from zipfile import ZipFile
    with ZipFile(io.BytesIO(data)) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml")
    root = ElementTree.fromstring(xml)
    payment = next(cell for cell in root.iter() if cell.attrib.get("r") == "L2")
    value = next(child for child in payment if child.tag.rsplit("}", 1)[-1] == "v")
    value.text = "1.00000000000000001"
    changed = replace_zip(data, "xl/worksheets/sheet1.xml", lambda _: ElementTree.tostring(root))
    assert mr_transfer.parse(changed, "exact.xlsx")[0]["values"]["paymentLimit"] == "invalid numeric payment"


def test_large_page_bulk_queries_exact_projection_pagination_and_ledger(client, record_property):
    from mr_projection_fixture import seed_directory, measured_queries
    api, db, _ = client
    headers, actor, _, _ = setup(api, db)
    ids = seed_directory(db, actor.id, 240)
    db.expunge_all()  # No warm identity-map objects can hide per-row lookups.
    with measured_queries(db.get_bind()) as metric:
        response = api.get(BASE, headers=headers, params={"limit": 100, "offset": 100, "query": "Bulk MR"})
    assert response.status_code == 200, response.text
    record_property("mr_page_queries", len(metric["statements"]))
    record_property("mr_page_seconds", metric["seconds"])
    assert len(metric["statements"]) <= 30
    assert max(metric["bind_counts"]) <= 500
    assert metric["seconds"] < 3
    result = response.json()
    assert (result["total"], result["filtered"], result["limit"], result["offset"]) == (240, 240, 100, 100)
    assert [row["id"] for row in result["items"]] == [str(key) for key in reversed(ids[40:140])]
    for item in result["items"]:
        # Legacy single-record projection remains an independent exact oracle.
        row = db.get(MRDirectory, uuid.UUID(item["id"]))
        expected = mrs.projection(db, row)
        from app.schemas.mrs import MRDirectoryResponse
        assert item == MRDirectoryResponse.model_validate(expected).model_dump(mode="json")
        assert item["createdBy"] == "Super Admin" and item["updatedBy"] == "Backend user"
    assert api.get(BASE, headers=headers, params={"offset": 240}).json()["items"] == []
    tombstone = api.get(BASE, headers=headers, params={"query": "FIX-00000"}).json()["items"][0]
    assert tombstone["reportingManagerName"] == "Deleted manager"
    assert tombstone["assignmentWarnings"] == ["Reporting manager was deleted. Explicitly replace it before saving."]
    # Filtered full-directory export still requires a durable ledger acknowledgement.
    response = api.get(BASE + "/export", headers={**headers, "X-Download-Initiation": str(uuid.uuid4())},
                       params={"format": "csv", "query": "Bulk MR", "status": "inactive"})
    assert response.status_code == 200 and response.headers["X-Download-Log"]
    from app.db.download_models import DownloadLog
    ledger = db.get(DownloadLog, uuid.UUID(response.headers["X-Download-Log"]))
    assert (ledger.source, ledger.kind, ledger.format, ledger.provenance) == ("mr", "export", "CSV", "server_prepared")
    exported = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert len(exported) == 121
    assert all(row[mr_transfer.HEADERS.index("Status")] == "inactive" for row in exported[1:])
