"""Isolated Doctor persistence/validation/transfer tests using landed MR identities."""
import csv
import io
import re
import uuid
from decimal import Decimal
import pytest
from openpyxl import load_workbook, Workbook
from pydantic import ValidationError
from sqlalchemy import func, select, text, insert
from app.db.doctor_models import DoctorDirectory
from app.db.mr_models import MRDirectory
from app.db.zone_models import Zone
from app.db.models import AuditEvent, MRProfile, User
from app.schemas.doctors import DoctorFields
from app.services import doctor_transfer
from test_sessions import client
from test_mrs import setup, fields as mr_fields, add as add_mr
from app.core.security import utcnow

BASE = "/api/v1/admin/doctors"


def setup_doctor(api, db):
    headers, actor, hq, zone = setup(api, db)
    mr = add_mr(api, headers, mr_fields(hq, zone))["record"]
    return headers, actor, mr, zone


def fields(mr, **changes):
    return {**dict(name="Synthetic Doctor", phone="9000000001", dialCountry="IN", alternatePhone="9000000002",
                   email="doctor@example.com", contactRequirement="required", dateOfJoining="2020-01-01",
                   registrationNumber="REG-01", qualification="MBBS", clinicName="Clinic", mrId=mr["id"],
                   invoiceType="gst", gstNumber="07ABCDE1234F1Z5", drugLicenceNumber="DL-001",
                   orderDiscount="12.35", daysLimit=45, paymentLimit="1234567890123.45", status="active",
                   pincode="110001", addressLine1="First address", addressLine2="Second address",
                   landmark="Landmark", country="India", state="Delhi", city="District"), **changes}


def add(api, headers, body):
    response = api.post(BASE, headers=headers, json=body)
    assert response.status_code == 201, response.text
    return response.json()


def edit(api, headers, record, body):
    return api.post(BASE + "/" + record["id"] + "/edit", headers=headers,
                    json={**body, "expected_version": record["version"]})


def csv_file(bodies, mr_reference="user:synthetic.mr", header=None):
    header = header or doctor_transfer.HEADERS
    stream = io.StringIO()
    writer = csv.writer(stream)
    writer.writerow(header)
    for body in bodies:
        values = dict(body, mrId=mr_reference, zoneName="Zone for MRs", verification="verified", dialCode="+91")
        writer.writerow([values[key] for key, title in doctor_transfer.COLUMNS if title in header])
    return stream.getvalue().encode()


def review(api, headers, data, filename="doctors.csv"):
    return api.post(BASE + "/import/review", headers={**headers, "Content-Type": "application/octet-stream"},
                    content=data, params={"filename": filename})


def commit(api, headers, data, digest, filename="doctors.csv"):
    return api.post(BASE + "/import/commit", headers={**headers, "Content-Type": "application/octet-stream"},
                    content=data, params={"filename": filename, "digest": digest, "confirm": True})


def test_all_fields_audit_persist_and_version_conflict(client):
    api, db, _ = client
    headers, actor, mr, zone = setup_doctor(api, db)
    body = fields(mr)
    record = add(api, headers, body)
    detail = api.get(BASE + "/" + record["id"], headers=headers).json()
    for key, value in body.items():
        assert detail[key] == value
    assert detail["mrName"] == mr["name"] and detail["zoneId"] == zone
    assert detail["createdBy"] == detail["updatedBy"] == "Super Admin"
    assert detail["verification"] == "unverified"
    assert "password" not in detail
    changed = edit(api, headers, record, fields(mr, name="Changed Doctor", phone="", email="", contactRequirement="optional"))
    assert changed.status_code == 200, changed.text
    assert changed.json()["version"] == 2 and changed.json()["createdAt"] == detail["createdAt"]
    stale = edit(api, headers, record, body)
    assert stale.status_code == 409
    db.expire_all()
    row = db.get(DoctorDirectory, uuid.UUID(record["id"]))
    assert row.name == "Changed Doctor" and row.paymentLimit == Decimal("1234567890123.45")
    assert db.scalar(select(AuditEvent.id).where(AuditEvent.resource_id == row.id, AuditEvent.action == "doctor_directory_edit"))


@pytest.mark.parametrize("changes", [
    {"contactRequirement": "required", "phone": ""}, {"alternatePhone": "123"}, {"dialCountry": "ZZ"},
    {"email": "bad"}, {"dateOfJoining": "2099-01-01"}, {"dateOfJoining": "2020-02-30"},
    {"invoiceType": "gst", "gstNumber": ""}, {"gstNumber": "bad"},
    {"orderDiscount": "100.01"}, {"orderDiscount": "1.001"}, {"paymentLimit": "1e3"},
    {"paymentLimit": "12345678901234.00"}, {"paymentLimit": "-1"},
    {"paymentLimit": 9999999999999.991}, {"paymentLimit": True}, {"daysLimit": 1.2},
    {"daysLimit": True}, {"daysLimit": 2147483648}, {"pincode": "ABC"},
    {"zoneId": str(uuid.uuid4())}, {"createdBy": "forged"}, {"password": "not-supported"},
])
def test_business_validation(changes):
    with pytest.raises(ValidationError):
        DoctorFields(**fields({"id": str(uuid.uuid4())}, **changes))


def test_country_phone_optional_and_blank_defaults():
    body = DoctorFields(**fields({"id": str(uuid.uuid4())}, country="United Arab Emirates", dialCountry="AE",
                                phone="123456789", alternatePhone="", pincode="AB-123", gstNumber="ABCDE",
                                dateOfJoining="", orderDiscount="", paymentLimit=""))
    assert body.paymentLimit == 0 and body.dateOfJoining is None


def test_duplicates_contacts_status_and_atomic_bulk(client):
    api, db, _ = client
    headers, _, mr, _ = setup_doctor(api, db)
    one = add(api, headers, fields(mr))
    two = add(api, headers, fields(mr, registrationNumber="REG-02", contactRequirement="optional", phone="", email=""))
    duplicate = api.post(BASE, headers=headers, json=fields(mr, registrationNumber="  reg-01 "))
    assert duplicate.status_code == 409
    invalid = api.post(BASE + "/" + two["id"] + "/contact", headers=headers,
                       json={"expected_version": two["version"], "contactRequirement": "required"})
    assert invalid.status_code == 422
    status = api.post(BASE + "/" + one["id"] + "/status", headers=headers,
                      json={"expected_version": 1, "status": "inactive"})
    assert status.status_code == 200
    selection = [{"id": row["id"], "expected_version": row["version"]} for row in (two, one)]
    conflict = api.post(BASE + "/bulk", headers=headers, json={"selected": selection, "operation": "verification", "verification": "verified"})
    assert conflict.status_code == 409
    assert api.get(BASE + "/" + two["id"], headers=headers).json()["version"] == 1
    selection[1]["expected_version"] = 2
    result = api.post(BASE + "/bulk", headers=headers, json={"selected": selection, "operation": "verification", "verification": "verified"})
    assert result.status_code == 200 and all(r["verification"] == "verified" for r in result.json())
    selection = [{"id": r["id"], "expected_version": r["version"]} for r in result.json()]
    missing = api.post(BASE + "/bulk", headers=headers, json={"selected": selection, "operation": "shift", "mrId": str(uuid.uuid4())})
    assert missing.status_code == 409
    assert api.get(BASE + "/" + two["id"], headers=headers).json()["mrId"] == mr["id"]


def test_assignment_retention_live_zone_and_missing_filters(client):
    api, db, _ = client
    headers, _, mr, zone = setup_doctor(api, db)
    one = add(api, headers, fields(mr))
    inactive = api.post("/api/v1/admin/mrs/" + mr["id"] + "/status", headers=headers,
                       json={"expected_version": mr["version"], "status": "inactive"})
    assert inactive.status_code == 200
    assert edit(api, headers, one, fields(mr, name="Retained inactive")).status_code == 200
    assert api.post(BASE, headers=headers, json=fields(mr, registrationNumber="NEW")).status_code == 409
    refs = api.get(BASE + "/references", headers=headers).json()
    assert refs["items"][0]["usable"] is False
    assert refs["items"][0]["zoneStatus"] == "active"
    # Deletion retains MR history/FK, but Doctor projections never pretend it is live.
    removed = api.post("/api/v1/admin/mrs/" + mr["id"] + "/delete", headers=headers,
                       json={"expected_version": inactive.json()["version"]})
    assert removed.status_code == 200
    missing = api.get(BASE, headers=headers, params={"mr_id": "missing"}).json()
    assert missing["filtered"] == 1 and missing["items"][0]["mrName"] == ""
    assert api.get(BASE + "/filters", headers=headers).json()["missingMR"]
    assert api.get(BASE + "/filters", headers=headers).json()["missingZone"]
    assert api.get(BASE, headers=headers, params={"zone_id": "missing"}).json()["filtered"] == 1
    assert api.get(BASE, headers=headers, params={"zone_id": zone}).json()["filtered"] == 0


def test_search_filter_pagination_export_parity_and_xlsx_text(client):
    api, db, _ = client
    headers, _, mr, zone = setup_doctor(api, db)
    for index in range(13):
        add(api, headers, fields(mr, name=f"=Doctor {index}" if index == 0 else f"Doctor {index}",
                                registrationNumber=f"00{index}", pincode="010001", status="inactive" if index % 2 else "active"))
    params = {"status": "active", "zone_id": zone, "mr_id": mr["id"], "state": "Delhi", "query": "Clinic"}
    data = api.get(BASE, headers=headers, params={**params, "limit": 2, "offset": 2}).json()
    assert data["filtered"] == 7 and data["total"] == 13 and len(data["items"]) == 2
    csv_response = api.get(BASE + "/export", headers=headers, params=params)
    assert csv_response.status_code == 200, csv_response.text
    rows = list(csv.reader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
    assert len(rows) == 8 and rows[0] == doctor_transfer.CURRENT_HEADERS + doctor_transfer.AUDIT
    assert all(row[24] == "'010001" for row in rows[1:])
    assert all(row[11] == "user:synthetic.mr" and row[28] == mr["name"] for row in rows[1:])
    assert any(row[0].startswith("'=") for row in rows[1:])
    assert csv_response.headers["x-download-log"]
    excel = api.get(BASE + "/export", headers=headers, params={**params, "format": "xlsx"})
    book = load_workbook(io.BytesIO(excel.content))
    assert book.active["Y2"].data_type == "s" and book.active["Y2"].value == "010001"
    assert book.active["S2"].data_type == "s"
    assert api.get(BASE, headers=headers, params={"query": "%_"}).json()["filtered"] == 0
    assert api.get(BASE + "/export", headers=headers, params={"query": "%_"}).status_code == 200
    assert api.get(BASE, headers=headers, params={"mr_id": "local-mr-id"}).status_code == 422


def test_review_commit_legacy_binding_and_rereview_conflict(client):
    api, db, _ = client
    headers, _, mr, _ = setup_doctor(api, db)
    body = fields(mr)
    data = csv_file([body], header=doctor_transfer.LEGACY)
    response = review(api, headers, data)
    assert response.status_code == 200 and response.json()["valid"], response.text
    assert db.scalar(select(func.count()).select_from(DoctorDirectory)) == 0
    assert commit(api, headers, data, response.json()["digest"], filename="renamed.csv").status_code == 409
    committed = commit(api, headers, data, response.json()["digest"])
    assert committed.status_code == 200 and committed.json()["imported"] == 1
    record = api.get(BASE, headers=headers).json()["items"][0]
    assert record["contactRequirement"] == "optional" and record["verification"] == "verified"
    assert record["createdBy"] == "Super Admin"
    assert commit(api, headers, data, response.json()["digest"]).status_code == 409
    exported = api.get(BASE + "/export", headers=headers).content
    assert review(api, headers, exported).json()["rows"][0]["errors"] == ["Registration number already belongs to another doctor."]
    new = csv_file([fields(mr, registrationNumber="NEW-01")])
    reviewed = review(api, headers, new).json()
    changed = api.post("/api/v1/admin/mrs/" + mr["id"] + "/contact", headers=headers,
                       json={"expected_version": mr["version"], "contactRequirement": "optional"})
    assert changed.status_code == 200
    assert commit(api, headers, new, reviewed["digest"]).status_code == 409


def test_import_atomic_invalid_duplicates_ambiguous_and_exact_xlsx(client):
    api, db, _ = client
    headers, _, mr, _ = setup_doctor(api, db)
    data = csv_file([fields(mr), fields(mr, registrationNumber="REG-02", invoiceType="gst", gstNumber="bad")])
    response = review(api, headers, data).json()
    assert not response["valid"] and not response["rows"][0]["errors"] and response["rows"][1]["errors"]
    assert commit(api, headers, data, response["digest"]).status_code == 409
    assert db.scalar(select(func.count()).select_from(DoctorDirectory)) == 0
    local = review(api, headers, csv_file([fields(mr)], mr_reference="sample-mr-1")).json()
    assert not local["valid"]
    # Duplicate labels need explicit account/employee/UUID references.
    add_mr(api, headers, {**mr_fields(mr["hq"], mr["zoneId"]), "employeeCode": "MR-02", "userId": "second.mr"})
    ambiguous = review(api, headers, csv_file([fields(mr)], mr_reference=mr["name"])).json()
    assert not ambiguous["valid"] and "ambiguous" in ambiguous["rows"][0]["errors"][0]
    book = Workbook()
    sheet = book.active
    sheet.append(doctor_transfer.HEADERS)
    values = {**fields(mr), "mrId": "user:synthetic.mr", "zoneName": "Zone for MRs", "dialCode": "+91", "verification": "unverified"}
    sheet.append([str(values[key]) for key, _ in doctor_transfer.COLUMNS])
    buffer = io.BytesIO(); book.save(buffer)
    response = review(api, headers, buffer.getvalue(), "doctors.xlsx")
    assert response.json()["valid"], response.text
    assert commit(api, headers, buffer.getvalue(), response.json()["digest"], "doctors.xlsx").json()["imported"] == 1
    assert api.get(BASE, headers=headers).json()["items"][0]["paymentLimit"] == "1234567890123.45"


def test_limits_bad_files_denials_and_unknown_fields(client):
    api, db, _ = client
    assert api.get(BASE).status_code == 401
    headers, _, mr, _ = setup_doctor(api, db)
    assert review(api, headers, b"x" * (2 * 1024 * 1024 + 1)).status_code == 413
    assert review(api, headers, csv_file([fields(mr)] * 1001)).status_code == 422
    assert review(api, headers, b"bad", "fake.xlsx").status_code == 422
    assert review(api, headers, b"Name,Password\nfake,secret").status_code == 422
    forged = api.post(BASE, headers=headers, json=fields(mr, **{"private-secret-field": "secret"}))
    assert forged.status_code == 422 and "private-secret-field" not in forged.text and "secret" not in forged.text
    # MR has no Admin grants; its valid portal token still cannot use Doctor services.
    account = add_mr(api, headers, {**mr_fields(mr["hq"], mr["zoneId"]), "employeeCode": "MR-03", "userId": "denied.mr"})
    login = api.post("/api/v1/auth/login", headers={"Origin": "http://testserver"}, json={
        "identifier": account["credentials"]["userId"], "password": account["credentials"]["password"], "identity_kind": "mr"})
    denied = {"Authorization": "Bearer " + login.json()["access_token"]}
    for path in ("", "/references", "/filters", "/sample", "/export"):
        assert api.get(BASE + path, headers=denied).status_code == 403


def test_migration_constraints_and_empty_directory(client):
    from pathlib import Path
    from alembic.config import Config
    from alembic.script import ScriptDirectory
    api, db, _ = client
    assert db.scalar(select(func.count()).select_from(DoctorDirectory)) == 0
    # This fixture upgrades to head, unlike explicitly historical migration
    # fixtures. Keep it current as unrelated master tables are added.
    heads = set(ScriptDirectory.from_config(Config(str(Path(__file__).resolve().parents[1] / "alembic.ini"))).get_heads())
    assert len(heads) == 1
    assert set(db.scalars(text("SELECT version_num FROM alembic_version"))) == heads
    constraints = set(db.scalars(text("SELECT conname FROM pg_constraint WHERE conrelid = 'doctor_directory'::regclass")))
    assert {"ck_doctor_version", "ck_doctor_limits", "ck_doctor_required_contact", "ck_doctor_invoice"} <= constraints
    assert db.scalar(text("SELECT count(*) FROM pg_indexes WHERE tablename='doctor_directory' AND indexname='uq_doctor_registration'")) == 1


def test_zone_grants_never_authorize_doctors(client):
    from test_zone_permissions import setup as staff_setup, staff_login, bearer
    api, db, _ = client
    _, _, _, record, password = staff_setup(api, db, grants=("zone.add", "zone.edit", "zone.delete", "zone.import", "zone.export"))
    headers = bearer(staff_login(api, record, password))
    for path in ("", "/references", "/filters", "/sample", "/export"):
        assert api.get(BASE + path, headers=headers).status_code == 403
    assert api.post(BASE, headers=headers, json=fields({"id": str(uuid.uuid4())})).status_code == 403
    assert review(api, headers, b"Name,MR\nBad,local").status_code == 403
    assert commit(api, headers, b"Name,MR\nBad,local", "0" * 64).status_code == 403


def test_complete_mr_choices_beyond_one_page_and_export_limit(client):
    api, db, _ = client
    headers, actor, mr, _ = setup_doctor(api, db)
    existing = db.get(MRDirectory, uuid.UUID(mr["id"]))
    prototype = {attribute.key: getattr(existing, attribute.key) for attribute in MRDirectory.__mapper__.column_attrs
                 if attribute.key not in ("id", "created_at", "updated_at", "deleted_at", "deleted_by")}
    template_user = db.get(User, db.get(MRProfile, existing.id).user_id)
    last_id = None
    for index in range(105):
        user = User(email=f"reference-{index}@example.com", username=f"reference.mr.{index}", password_hash=template_user.password_hash,
                    system_role="mr", identity_version=1)
        db.add(user); db.flush()
        profile = MRProfile(user_id=user.id, is_active=True)
        db.add(profile); db.flush()
        db.add(MRDirectory(**{**prototype, "name": f"ZZ Reference MR {index:03}", "employeeCode": f"REF-{index}"}, id=profile.id))
        last_id = str(profile.id)
    db.commit()
    first = api.get(BASE + "/references", headers=headers, params={"limit": 100}).json()
    second = api.get(BASE + "/references", headers=headers, params={"limit": 100, "offset": 100}).json()
    assert first["total"] == second["total"] == 106
    assert len(first["items"]) == 100 and len(second["items"]) == 6
    assert last_id in {row["id"] for row in second["items"]}
    values = DoctorFields(**fields(mr)).model_dump()
    db.execute(insert(DoctorDirectory), [{**values, "registrationNumber": f"LIMIT-{index}",
                                         "created_by": actor.id, "updated_by": actor.id, "verification": "unverified",
                                         "created_at": utcnow(), "updated_at": utcnow()}
                                        for index in range(5001)])
    db.commit()
    response = api.get(BASE + "/export", headers=headers)
    assert response.status_code == 409 and "5,000" in response.text
    assert api.get(BASE, headers=headers, params={"limit": 1, "offset": 5000}).json()["filtered"] == 5001


def test_atomic_shift_and_zone_reassignment_are_live_not_snapshot(client):
    api, db, _ = client
    headers, _, mr, _ = setup_doctor(api, db)
    one = add(api, headers, fields(mr))
    two = add(api, headers, fields(mr, registrationNumber="SECOND"))
    zone = api.post("/api/v1/admin/zones", headers=headers, json={"name": "New Doctor Zone", "status": "active"}).json()
    target = add_mr(api, headers, {**mr_fields(mr["hq"], zone["id"]), "employeeCode": "TARGET", "userId": "target.mr"})["record"]
    response = api.post(BASE + "/bulk", headers=headers, json={"operation": "shift", "mrId": target["id"],
                      "selected": [{"id": row["id"], "expected_version": row["version"]} for row in (one, two)]})
    assert response.status_code == 200 and all(row["zoneId"] == zone["id"] and row["mrId"] == target["id"] for row in response.json())
    third = api.post("/api/v1/admin/zones", headers=headers, json={"name": "Later MR Zone", "status": "active"}).json()
    changed = api.post("/api/v1/admin/mrs/" + target["id"] + "/edit", headers=headers,
                      json={**mr_fields(mr["hq"], third["id"]), "employeeCode": target["employeeCode"], "userId": target["userId"], "expected_version": target["version"]})
    assert changed.status_code == 200, changed.text
    assert api.get(BASE + "/" + one["id"], headers=headers).json()["zoneId"] == third["id"]
    assert api.get(BASE, headers=headers, params={"zone_id": zone["id"]}).json()["filtered"] == 0
    assert api.get(BASE, headers=headers, params={"zone_id": third["id"]}).json()["filtered"] == 2


def test_xlsx_uses_original_numeric_lexemes_and_rejects_formulas(client):
    import zipfile
    api, db, _ = client
    headers, _, mr, _ = setup_doctor(api, db)
    values = {**fields(mr), "mrId": "user:synthetic.mr", "zoneName": "", "dialCode": "+91", "verification": "unverified"}
    book = Workbook(); sheet = book.active
    sheet.append(doctor_transfer.HEADERS)
    sheet.append([str(values[key]) for key, _ in doctor_transfer.COLUMNS])
    sheet["S2"] = 9999999999999.99
    buffer = io.BytesIO(); book.save(buffer)
    def alter(xml):
        source = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
        output = io.BytesIO()
        with source, zipfile.ZipFile(output, "w") as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    data = xml(data.decode()).encode()
                target.writestr(item, data)
        return output.getvalue()
    oversized_precision = alter(lambda xml: re.sub(r'(<c r="S2"[^>]*><v>)[^<]+', r'\g<1>9999999999999.991', xml))
    report = review(api, headers, oversized_precision, "precision.xlsx")
    assert report.status_code == 200 and not report.json()["valid"], report.text
    formula = alter(lambda xml: xml.replace('<c r="S2" t="n"><v>', '<c r="S2" t="n"><f>1+1</f><v>'))
    assert review(api, headers, formula, "formula.xlsx").status_code == 422
